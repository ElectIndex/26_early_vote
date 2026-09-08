"""Connecticut: the SOTS per-ballot early-voting and absentee detail files.

Two real fixtures, both from `portal.ct.gov`'s 2026 Voter Data page as it stood
on 2026-09-06:

* `2026-09-03_early_voting_sample.xlsx` -- one row per in-person early vote.
* `2026-09-03_absentee_ballot_sample.xlsx` -- one row per absentee ballot
  ISSUED, with a blank `DT RETURNED` for the ones still out.

Both carry Connecticut's 33-column header verbatim and every value in every
column this adapter reads. The columns it does NOT read -- name, street address,
mailing address, ZIP, year of birth, voter and ballot serial numbers -- are
blanked, because a fixture lives in a public repository and republishing sixty
real voters' home addresses to catch a schema change is not a trade worth making.
The `VOTER ID` column is renumbered 1..n so the shape of the column survives.

These files are the September 1, 2026 special DEMOCRATIC primary in Enfield,
which is the only election Connecticut had files up for. That has two useful
consequences for the tests: they are the exact case the "which election is this?"
guard exists to refuse, and they cannot show what a general-election file's party
vocabulary looks like -- which is why `PARTY_ALIASES` is documented as
unverified and why an unknown label raises SchemaDrift.
"""

from __future__ import annotations

import io
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

from ev.adapters import _towns, ct
from ev.adapters._net import Missing
from ev.adapters.base import NotYetPublished, SchemaDrift, SourceError

FIXTURES = Path(__file__).parent / "fixtures" / "ct"
EARLY = FIXTURES / "2026-09-03_early_voting_sample.xlsx"
ABSENTEE = FIXTURES / "2026-09-03_absentee_ballot_sample.xlsx"
INDEX = FIXTURES / "2026-voter-data-index-excerpt.html"

ENFIELD = "0900325990"          # Enfield town, Hartford County
HARTFORD_COUNTY = "09003"
GREENWICH = "0900133620"        # Greenwich town, Fairfield County
FAIRFIELD_COUNTY = "09001"

#: The fixtures' ballots run 2026-08-14..2026-09-01, which is 63-81 days before
#: the 2026 general. Shifting them forward by this many days lands them on
#: 2026-10-01..2026-10-19 -- inside the general's own window -- without touching
#: a single count.
SHIFT = 48


def _edit(path: Path, mutate) -> bytes:
    """Round-trip a fixture through `mutate(header, row)` and return xlsx bytes."""
    book = openpyxl.load_workbook(path)
    sheet = book.worksheets[0]
    rows = list(sheet.iter_rows())
    header = [" ".join(str(c.value or "").split()).upper() for c in rows[0]]
    for row in rows[1:]:
        mutate(header, row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _shift_dates(days: int = SHIFT):
    def mutate(header, row):
        for column in ("DT MAILED", "DT RETURNED"):
            cell = row[header.index(column)]
            if cell.value:
                cell.value = (date.fromisoformat(str(cell.value)[:10])
                              + timedelta(days=days)).isoformat()
    return mutate


def _set(column: str, value, *, where=lambda header, row: True):
    def mutate(header, row):
        if where(header, row):
            row[header.index(column)].value = value
    return mutate


def _shifted(path: Path) -> bytes:
    return _edit(path, _shift_dates())


@pytest.fixture(scope="module")
def raw() -> ct.Tally:
    tally = ct.Tally()
    ct.read(ABSENTEE.read_bytes(), "mail", tally)
    ct.read(EARLY.read_bytes(), "inperson", tally)
    return tally


@pytest.fixture(scope="module")
def shifted() -> list[tuple[str, bytes]]:
    return [("mail", _shifted(ABSENTEE)), ("inperson", _shifted(EARLY))]


@pytest.fixture(scope="module")
def result(shifted):
    return ct.build(shifted, 2026, date(2026, 9, 3) + timedelta(days=SHIFT))


# --------------------------------------------------------------------------
# Reading the two files
# --------------------------------------------------------------------------
def test_the_two_files_are_the_method_split(raw):
    # 30 absentee ballots back and 28 early votes -- from two files, with no
    # method label interpreted anywhere.
    running = ct._Bucket()
    for _, bucket in sorted(raw.by_day.items()):
        running.add(bucket)
    assert (running.returned, running.inperson) == (30, 28)
    assert running.total == raw.counted == 58


def test_absentee_ballots_are_counted_as_issued_on_the_day_they_were_mailed(raw):
    running = ct._Bucket()
    for _, bucket in sorted(raw.by_day.items()):
        running.add(bucket)
    # 36 issued, of which 30 are back: the six still out are real and are on the
    # issued curve without being on the returns curve.
    assert running.issued == 36
    assert running.issued > running.returned


def test_void_ballots_are_dropped_from_every_count():
    """The fixtures each carry one real `RETURN_TYPE == "Void"` row."""
    kept = ct.Tally()
    ct.read(EARLY.read_bytes(), "inperson", kept)
    voided = ct.Tally()
    ct.read(_edit(EARLY, _set("RETURN_TYPE", "In Person By Voter")), "inperson", voided)
    assert voided.counted == kept.counted + 1


def test_an_unfamiliar_return_type_is_not_drift():
    # RETURN_TYPE is read for the word "Void" and nothing else, so a new drop
    # box or delivery route must not break Connecticut.
    tally = ct.Tally()
    ct.read(_edit(EARLY, _set("RETURN_TYPE", "Drop Box 97")), "inperson", tally)
    assert tally.counted == 29


def test_the_town_is_resolved_to_a_census_geoid(raw):
    assert raw.town_names == {ENFIELD: "Enfield town"}
    assert not raw.unmapped


def test_party_comes_off_every_ballot(raw):
    assert raw.unenrolled == 0
    running = ct._Bucket()
    for _, bucket in sorted(raw.by_day.items()):
        running.add(bucket)
    assert running.party == {"dem": 58}


# --------------------------------------------------------------------------
# Which election is this?
# --------------------------------------------------------------------------
def test_a_special_primarys_ballots_are_refused_for_the_general(raw):
    # The fixtures ARE the September 1 special primary, sitting in the same
    # folder the general's files will use. Reading them as general-election
    # early voting is the mistake this guard exists to prevent.
    assert ct.belongs_to(raw, 2026) is False
    with pytest.raises(NotYetPublished) as caught:
        ct.build([("mail", ABSENTEE.read_bytes()), ("inperson", EARLY.read_bytes())],
                 2026, date(2026, 9, 3))
    assert "another election" in str(caught.value)


def test_ballots_inside_the_generals_window_are_accepted(shifted):
    tally = ct.Tally()
    for kind, body in shifted:
        ct.read(body, kind, tally)
    assert ct.belongs_to(tally, 2026) is True


def test_the_filename_window_opens_45_days_out():
    assert ct.in_file_window(date(2026, 9, 6), 2026) is False
    assert ct.in_file_window(date(2026, 9, 20), 2026) is True
    assert ct.in_file_window(date(2026, 11, 3), 2026) is True
    assert ct.in_file_window(date(2026, 12, 1), 2026) is False


# --------------------------------------------------------------------------
# The rows we publish
# --------------------------------------------------------------------------
def test_state_rows_are_a_cumulative_daily_curve(result):
    totals = [row.ballots_total for row in result.state_rows]
    assert totals == sorted(totals)
    assert totals[-1] == 58
    last = result.state_rows[-1]
    assert last.day == date(2026, 9, 3) + timedelta(days=SHIFT)
    assert (last.mail_requested, last.mail_returned, last.inperson) == (36, 30, 28)
    assert last.mail_returned + last.inperson == last.ballots_total


def test_party_is_published_as_real_counts(result):
    last = result.state_rows[-1]
    assert last.party_dem == 58
    # Connecticut enrols by party and this file names it on every ballot, so a
    # party with no ballots yet is a real 0 -- blank would say Connecticut does
    # not report party at all.
    assert (last.party_rep, last.party_npa, last.party_oth) == (0, 0, 0)


def test_town_rows_are_keyed_by_geoid(result):
    towns = _towns.rows_of(result)
    assert towns, "town rows are the point of a New England adapter"
    assert {row.town_geoid for row in towns} == {ENFIELD}
    assert towns[-1].town_name == "Enfield town"
    assert towns[-1].ballots_total == 58


def test_counties_fall_out_of_the_geoid_and_sum_exactly(result):
    counties = {row.county_fips for row in result.county_rows}
    assert counties == {HARTFORD_COUNTY}
    last = result.county_rows[-1]
    assert last.county_name == "Hartford County"
    assert last.ballots_total == result.state_rows[-1].ballots_total


def test_two_towns_in_two_counties_roll_up_to_two_counties():
    """Rename half the ballots to a Fairfield County town and the county table
    has to split, with no crosswalk anywhere in the code."""
    def move(header, row):
        town = header.index("RESIDENCE ADDRESS CITY")
        if int(row[header.index("VOTER ID")].value) % 2 == 0:
            row[town].value = "Greenwich"
    bodies = [("mail", _edit(ABSENTEE, lambda h, r: (_shift_dates()(h, r), move(h, r)))),
              ("inperson", _edit(EARLY, lambda h, r: (_shift_dates()(h, r), move(h, r))))]
    result = ct.build(bodies, 2026, date(2026, 9, 3) + timedelta(days=SHIFT))
    assert {row.town_geoid for row in _towns.rows_of(result)} == {ENFIELD, GREENWICH}
    finals = {row.county_fips: row for row in result.county_rows
              if row.day == max(r.day for r in result.county_rows)}
    assert set(finals) == {HARTFORD_COUNTY, FAIRFIELD_COUNTY}
    assert sum(row.ballots_total for row in finals.values()) == 58
    assert finals[FAIRFIELD_COUNTY].county_name == "Fairfield County"


def test_a_file_with_only_absentee_leaves_inperson_blank(shifted):
    mail_only = [entry for entry in shifted if entry[0] == "mail"]
    result = ct.build(mail_only, 2026, date(2026, 9, 3) + timedelta(days=SHIFT))
    last = result.state_rows[-1]
    # Early voting has not opened, so Connecticut publishes no early-voting
    # file. That is "not reported", not "nobody voted early". See THE BLANK RULE.
    assert last.inperson is None
    assert last.mail_returned == 30
    assert last.ballots_total == 30
    assert _towns.rows_of(result)[-1].inperson is None


# --------------------------------------------------------------------------
# ⚠️ A CUMULATIVE SERIES MAY NEVER FALL
#
# `_Bucket.total` is `returned + inperson`, and an UNREAD method contributes 0
# rather than nothing. On these fixtures both files give 58, the absentee file
# alone gives 30 and the early-voting file alone gives 28 -- so a run that lost
# one workbook published a cumulative total below yesterday's. These tests are
# that failure, in both directions, and the two guards that close it.
# --------------------------------------------------------------------------
def test_early_voting_alone_is_refused_rather_than_published_as_a_total(shifted):
    """58 -> 28 was the measured fall. Connecticut issues absentee ballots two
    weeks before early voting opens, so an early-voting file with no absentee
    partner is a failed fetch and must fall through, not publish."""
    early_only = [entry for entry in shifted if entry[0] == "inperson"]
    with pytest.raises(SourceError) as caught:
        ct.build(early_only, 2026, date(2026, 9, 3) + timedelta(days=SHIFT))
    assert "absentee file was not" in str(caught.value)


def test_require_both_is_a_pure_check_on_which_files_were_read():
    ct.require_both({"mail"})                 # the opening fortnight: legitimate
    ct.require_both({"mail", "inperson"})     # the normal case
    ct.require_both(set())                    # nothing read; the caller's problem
    with pytest.raises(SourceError):
        ct.require_both({"inperson"})


def test_the_published_total_is_monotone_across_the_two_phases(shifted):
    """Absentee-only, then both: the only two shapes that can now be published,
    and the second can never be smaller than the first."""
    through = date(2026, 9, 3) + timedelta(days=SHIFT)
    mail_only = ct.build([e for e in shifted if e[0] == "mail"], 2026, through)
    both = ct.build(list(shifted), 2026, through)
    assert mail_only.state_rows[-1].ballots_total == 30
    assert both.state_rows[-1].ballots_total == 58
    assert both.state_rows[-1].ballots_total >= mail_only.state_rows[-1].ballots_total
    for result in (mail_only, both):
        totals = [row.ballots_total for row in result.state_rows]
        assert totals == sorted(totals)


# --------------------------------------------------------------------------
# Drift
# --------------------------------------------------------------------------
def test_a_renamed_town_column_is_drift():
    body = _edit(EARLY, lambda h, r: None)
    book = openpyxl.load_workbook(io.BytesIO(body))
    sheet = book.worksheets[0]
    for cell in sheet[1]:
        if " ".join(str(cell.value or "").split()).upper() == "RESIDENCE ADDRESS CITY":
            cell.value = "RES CITY"
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(SchemaDrift) as caught:
        ct.read(buffer.getvalue(), "inperson", ct.Tally())
    assert "town" in str(caught.value)


def test_an_unrecognised_party_is_drift():
    body = _edit(EARLY, _set("PARTY", "Whig"))
    with pytest.raises(SchemaDrift) as caught:
        ct.read(body, "inperson", ct.Tally())
    assert "Whig" in str(caught.value)


def test_connecticuts_independent_party_is_a_party_not_an_unaffiliated_voter():
    # normalize.party("independent") returns "npa", which is right almost
    # everywhere and wrong in Connecticut, where the Independent Party is a
    # registered party and unaffiliated voters are called "Unaffiliated".
    assert ct._party("Independent") == "oth"
    assert ct._party("Independent Party") == "oth"
    assert ct._party("Unaffiliated") == "npa"
    assert ct._party("Democratic") == "dem"
    assert ct._party("Republican") == "rep"
    assert ct._party("") is None


def test_a_town_we_cannot_place_is_excluded_counted_and_eventually_drift():
    def rename(header, row):
        _shift_dates()(header, row)
        if int(row[header.index("VOTER ID")].value) % 2 == 0:
            row[header.index("RESIDENCE ADDRESS CITY")].value = "Thompsonville"

    tally = ct.Tally()
    ct.read(_edit(ABSENTEE, rename), "mail", tally)
    # Counted and named, never guessed into a county.
    assert tally.unmapped["Thompsonville"] > 0
    assert ENFIELD in tally.town_names
    # ... and well past the floor, so the whole file is refused rather than
    # publishing half a map.
    with pytest.raises(SchemaDrift) as caught:
        ct._report_coverage(tally)
    assert "county-subdivision GEOID" in str(caught.value)


def test_blank_party_cells_suppress_the_party_columns_rather_than_understate():
    def blank(header, row):
        _shift_dates()(header, row)
        if int(row[header.index("VOTER ID")].value) % 2 == 0:
            row[header.index("PARTY")].value = None

    bodies = [("mail", _edit(ABSENTEE, blank)), ("inperson", _edit(EARLY, blank))]
    result = ct.build(bodies, 2026, date(2026, 9, 3) + timedelta(days=SHIFT))
    last = result.state_rows[-1]
    assert last.ballots_total == 58
    assert (last.party_dem, last.party_rep, last.party_npa, last.party_oth) == (None,) * 4
    assert _towns.rows_of(result)[-1].party_dem is None


def test_a_header_with_no_rows_is_drift():
    book = openpyxl.load_workbook(EARLY)
    sheet = book.worksheets[0]
    sheet.delete_rows(2, sheet.max_row)
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(SchemaDrift):
        ct.read(buffer.getvalue(), "inperson", ct.Tally())


def test_html_instead_of_a_workbook_is_a_source_error():
    with pytest.raises(SourceError):
        ct.read(b"<!doctype html><html><title>404 Error Page</title>", "mail", ct.Tally())


# --------------------------------------------------------------------------
# Finding the files
# --------------------------------------------------------------------------
def test_index_page_yields_one_dated_url_per_kind():
    found = ct.index_files(INDEX.read_bytes(), 2026)
    assert found["inperson"] == [(
        date(2026, 9, 3),
        "https://portal.ct.gov/-/media/sots/electionservices/2026_absentee_ballot_data/"
        "early_voting_09032026.xlsx")]
    assert found["mail"] == [(
        date(2026, 9, 3),
        "https://portal.ct.gov/-/media/sots/electionservices/2026_absentee_ballot_data/"
        "absentee_ballot_09032026.xlsx")]


def test_both_of_connecticuts_filename_conventions_are_read():
    # Verified in the wild inside one cycle: `abs_detail_073126.xlsx` on the
    # 2026-08-01 capture of the page, `absentee_ballot_09032026.xlsx` today.
    page = (b'<a href="/-/media/sots/electionservices/2026_absentee_ballot_data/'
            b'abs_detail_073126.xlsx?rev=abc&hash=def">Absentee Ballot Data</a>')
    found = ct.index_files(page, 2026)
    assert found["mail"][0][0] == date(2026, 7, 31)
    # The Sitecore rev/hash query is dropped: the bare path serves the same bytes.
    assert found["mail"][0][1].endswith("abs_detail_073126.xlsx")


def test_index_ignores_files_that_are_neither_kind():
    page = b'<a href="/-/media/sots/electionservices/2026_absentee_ballot_data/nov26re_11032026.xlsx">x</a>'
    assert ct.index_files(page, 2026) == {"inperson": [], "mail": []}


def test_portal_ct_gov_answers_a_missing_page_with_a_200_and_that_is_not_absence(monkeypatch):
    soft = b"<html><head><title>404 Error Page</title></head><body></body></html>"
    monkeypatch.setattr(ct, "get", lambda url, **kw: soft)
    # "We could not look" must fall through to a weaker tier, never stop the
    # ladder the way "there is no data yet" does.
    with pytest.raises(SourceError):
        ct.CTScraper().fetch(2026, date(2026, 10, 20))


def test_fetch_before_the_window_opens_is_not_yet_published(monkeypatch):
    def explode(*args, **kwargs):  # pragma: no cover - must never be called
        raise AssertionError("no request should be made outside the window")

    monkeypatch.setattr(ct, "get", explode)
    with pytest.raises(NotYetPublished):
        ct.CTScraper().fetch(2026, date(2026, 9, 6))


def test_fetch_reads_both_files_and_dates_the_series_by_the_older(monkeypatch, shifted):
    bodies = {
        "voter-data": INDEX.read_bytes(),
        "early_voting_10212026.xlsx": dict(shifted)["inperson"],
        "absentee_ballot_10202026.xlsx": dict(shifted)["mail"],
    }

    def fake_get(url, *, state, filename, **kwargs):
        for marker, body in bodies.items():
            if marker in url:
                return body
        raise Missing(f"{state}: {url} returned 404")

    monkeypatch.setattr(ct, "get", fake_get)
    result = ct.CTScraper().fetch(2026, date(2026, 10, 21))
    # The absentee file is a day older, so the curve stops where BOTH files can
    # speak -- otherwise 10-21 would publish early votes against frozen mail.
    assert result.state_rows[-1].day == date(2026, 10, 20)
    assert result.state_rows[-1].ballots_total == 58
    assert all(row.provenance is not None for row in _towns.rows_of(result))


def test_fetch_with_no_file_in_the_lookback_is_not_yet_published(monkeypatch):
    def fake_get(url, *, state, filename, **kwargs):
        if "voter-data" in url:
            return b"<html><body>the page exists but links nothing</body></html>"
        raise Missing(f"{state}: {url} returned 404")

    monkeypatch.setattr(ct, "get", fake_get)
    with pytest.raises(NotYetPublished):
        ct.CTScraper().fetch(2026, date(2026, 10, 20))


def test_there_is_no_archive_to_backfill():
    with pytest.raises(NotYetPublished) as caught:
        ct.CTScraper().fetch_history(2024)
    assert "no 2024 archive" in str(caught.value)


def _page(*names: str) -> bytes:
    """A Voter Data page linking exactly these filenames."""
    return b"".join(
        b'<a href="/-/media/sots/electionservices/2026_absentee_ballot_data/'
        + name.encode() + b'">x</a>'
        for name in names
    )


def test_a_listed_file_that_will_not_download_is_a_failure_not_an_absence(monkeypatch, shifted):
    """The mirror of `require_both`: Connecticut's own page says an
    early-voting file exists for this election, and we could not get it. Reading
    that as "no early voting today" would publish an absentee-only total -- 58
    back down to 30 -- so it falls through instead."""
    page = _page("absentee_ballot_10202026.xlsx", "early_voting_10202026.xlsx")

    def fake_get(url, *, state, filename, **kwargs):
        if "voter-data" in url:
            return page
        if "absentee_ballot_10202026" in url:
            return dict(shifted)["mail"]
        raise Missing(f"{state}: {url} returned 404")

    monkeypatch.setattr(ct, "get", fake_get)
    with pytest.raises(SourceError) as caught:
        ct.CTScraper().fetch(2026, date(2026, 10, 20))
    assert "inperson" in str(caught.value)


def test_a_kind_the_page_does_not_list_at_all_is_the_legitimate_opening_phase(monkeypatch, shifted):
    """Before early voting opens there is no early-voting file and the page
    lists none, so the absentee-only total is complete and is published."""
    page = _page("absentee_ballot_10202026.xlsx")

    def fake_get(url, *, state, filename, **kwargs):
        if "voter-data" in url:
            return page
        if "absentee_ballot_10202026" in url:
            return dict(shifted)["mail"]
        raise Missing(f"{state}: {url} returned 404")

    monkeypatch.setattr(ct, "get", fake_get)
    result = ct.CTScraper().fetch(2026, date(2026, 10, 20))
    last = result.state_rows[-1]
    assert last.ballots_total == 30
    assert last.mail_returned == 30
    # Never 0: Connecticut has not published an early-voting file at all.
    assert last.inperson is None


def test_a_listed_file_outside_the_generals_window_is_not_expected(monkeypatch, shifted):
    """The September special primary's files sit in the same folder. They are
    listed, they are not this election's, and they must not make a run that
    ignores them look like a failed fetch."""
    page = _page("absentee_ballot_10202026.xlsx", "early_voting_09032026.xlsx")

    def fake_get(url, *, state, filename, **kwargs):
        if "voter-data" in url:
            return page
        if "absentee_ballot_10202026" in url:
            return dict(shifted)["mail"]
        raise Missing(f"{state}: {url} returned 404")

    monkeypatch.setattr(ct, "get", fake_get)
    result = ct.CTScraper().fetch(2026, date(2026, 10, 20))
    assert result.state_rows[-1].ballots_total == 30
