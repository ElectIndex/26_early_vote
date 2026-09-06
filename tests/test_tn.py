"""Tennessee: the SoS early-voting and absentee workbook.

Two fixtures, because Tennessee shipped two different sheet layouts in two
consecutive cycles and both have to parse:

* `20241105EarlyAbsentee1031.xlsx` — `County | Date | EarlyVoting | Absentee |
  Day Total`, the 2024 general with a full method split.
* `20221108EarlyAbsenteethrough1103.xlsx` — `CoID | County | Date | Day Total`,
  the 2022 general, which reported **no** method split at all. Its mail and
  in-person fields must come out blank, not zero.

Both keep the real header verbatim and every row for Statewide plus five
counties. Expected numbers were accumulated straight from the fixtures with a
plain openpyxl loop.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from ev.adapters import tn
from ev.adapters._net import Missing
from ev.adapters.base import NotYetPublished, SchemaDrift, SourceError

FIXTURES = Path(__file__).parent / "fixtures" / "tn"
GEN24 = FIXTURES / "20241105EarlyAbsentee1031.xlsx"
GEN22 = FIXTURES / "20221108EarlyAbsenteethrough1103.xlsx"

LAST24 = date(2024, 10, 31)
LAST22 = date(2022, 11, 3)


@pytest.fixture(scope="module")
def gen2024():
    return tn.parse(GEN24.read_bytes(), 2024, LAST24)


@pytest.fixture(scope="module")
def gen2022():
    return tn.parse(GEN22.read_bytes(), 2022, LAST22)


def _rewrite(path: Path, edit) -> bytes:
    book = openpyxl.load_workbook(path)
    edit(book.worksheets[0])
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# THE BLANK RULE
# --------------------------------------------------------------------------
def test_no_party_field_is_ever_populated(gen2024, gen2022):
    """Tennessee has no party registration. The party columns in its PRIMARY
    files record which ballot a voter pulled, which is not the same thing and is
    not in these files at all -- a 0 here would read as 'zero Democrats voted'."""
    for result in (gen2024, gen2022):
        for row in result.state_rows + result.county_rows:
            for field in ("party_dem", "party_rep", "party_oth", "party_npa"):
                assert getattr(row, field) is None, (field, row)


def test_mail_requested_is_always_blank(gen2024):
    """Tennessee reports ballots RETURNED and never ballots requested."""
    assert all(r.mail_requested is None for r in gen2024.state_rows)


def test_2022_layout_reports_no_method_split_and_says_so(gen2022):
    """The 2022 sheet has no EarlyVoting/Absentee columns. Publishing 0 would
    claim Tennessee counted no absentee ballots that year."""
    for row in gen2022.state_rows + gen2022.county_rows:
        assert row.mail_returned is None, row
        assert row.inperson is None, row
        assert row.ballots_total is not None and row.ballots_total > 0


# --------------------------------------------------------------------------
# Daily counts accumulated into a curve
# --------------------------------------------------------------------------
def test_statewide_series_is_cumulative(gen2024):
    days = [r.day for r in gen2024.state_rows]
    assert days[0] == date(2024, 10, 16)
    assert days[-1] == LAST24
    assert days == sorted(days)
    assert len(days) == 14, "Tennessee skips Sundays; the file has 14 dates"

    totals = [r.ballots_total for r in gen2024.state_rows]
    assert totals[0] == 215936
    assert totals[4] == 820564
    assert totals[-1] == 2214870          # Tennessee's published 2024 EV total
    assert totals == sorted(totals)

    first, last = gen2024.state_rows[0], gen2024.state_rows[-1]
    assert first.ballots_new == 215936
    assert last.ballots_new == 177430
    assert last.inperson == 2132610
    assert last.mail_returned == 82253


def test_ballots_total_is_tennessees_own_day_total_not_our_sum(gen2024):
    """`Day Total` runs one to three ballots ahead of `EarlyVoting + Absentee`
    on nine rows of the real 2024 file. Deriving the headline from the split
    would leave us a few ballots off the Secretary of State's own number."""
    last = gen2024.state_rows[-1]
    assert last.inperson + last.mail_returned == 2214863
    assert last.ballots_total == 2214870, "Tennessee's figure, not ours"
    assert last.inperson + last.mail_returned < last.ballots_total

    oct28 = next(r for r in gen2024.state_rows if r.day == date(2024, 10, 28))
    assert oct28.ballots_new == 157210          # Day Total
    # ...against EarlyVoting 153,701 + Absentee 3,506 = 157,207.
    for row in gen2024.state_rows:
        assert (row.inperson or 0) + (row.mail_returned or 0) <= row.ballots_total


def test_2022_statewide_series(gen2022):
    rows = gen2022.state_rows
    assert rows[0].day == date(2022, 10, 19)
    assert rows[-1].day == LAST22
    assert rows[-1].ballots_total == 882285
    assert rows[0].ballots_new == 59178
    assert rows[-1].ballots_new == 111324


def test_as_of_truncates_the_series():
    """The adapter runs daily; a file posted on the 31st must not publish days
    the run has not reached."""
    partial = tn.parse(GEN24.read_bytes(), 2024, date(2024, 10, 19))
    assert [r.day for r in partial.state_rows] == [
        date(2024, 10, 16), date(2024, 10, 17),
        date(2024, 10, 18), date(2024, 10, 19),
    ]
    assert partial.state_rows[-1].ballots_total == 654223


def test_a_run_before_early_voting_opens_is_not_yet_published():
    """This is the path that runs every day for weeks. It must STOP the ladder,
    not fall through to a source that would invent a zero."""
    with pytest.raises(NotYetPublished, match="early voting opens"):
        tn.parse(GEN24.read_bytes(), 2024, date(2024, 9, 5))


# --------------------------------------------------------------------------
# Statewide is Tennessee's own row, not our sum
# --------------------------------------------------------------------------
def test_statewide_is_not_our_sum_of_the_counties(gen2024):
    """The fixture keeps five counties out of 95, so a parser that summed them
    would publish a fraction of Tennessee's own number."""
    final_counties = sum(
        r.ballots_total for r in gen2024.county_rows if r.day == LAST24
    )
    assert final_counties < gen2024.state_rows[-1].ballots_total
    assert gen2024.state_rows[-1].ballots_total == 2214870


def test_statewide_is_not_emitted_as_a_county(gen2024, gen2022):
    for result in (gen2024, gen2022):
        assert "Statewide" not in {r.county_name for r in result.county_rows}
        assert all(r.county_fips != "47000" for r in result.county_rows)


# --------------------------------------------------------------------------
# County rows are keyed by FIPS, never by name
# --------------------------------------------------------------------------
def test_county_rows_are_fips_keyed(gen2024):
    final = {r.county_fips: r for r in gen2024.county_rows if r.day == LAST24}
    assert final["47001"].county_name == "Anderson County"
    assert final["47041"].county_name == "DeKalb County"
    assert final["47157"].county_name == "Shelby County"
    assert final["47001"].ballots_total == 27187
    assert final["47001"].inperson == 26106
    assert final["47001"].mail_returned == 1080
    assert final["47041"].ballots_total == 6397
    assert final["47157"].ballots_total == 257733


def test_2022_counties_key_the_same_way(gen2022):
    final = {r.county_fips: r for r in gen2022.county_rows if r.day == LAST22}
    assert final["47001"].ballots_total == 11731
    assert final["47157"].ballots_total == 120651


def test_all_county_fips_are_five_digit_tennessee(gen2024, gen2022):
    for result in (gen2024, gen2022):
        for row in result.county_rows:
            assert len(row.county_fips) == 5 and row.county_fips.startswith("47")


# --------------------------------------------------------------------------
# Two layouts, and only two
# --------------------------------------------------------------------------
def test_both_shipped_layouts_are_allowed():
    assert frozenset({"county", "date", "earlyvoting", "absentee", "day total"}) in tn.LAYOUTS
    assert frozenset({"coid", "county", "date", "day total"}) in tn.LAYOUTS


def test_an_unlisted_layout_is_drift():
    """The failure this guards against is Tennessee renaming `EarlyVoting` and
    us silently publishing 'method not reported' for a cycle where they did."""
    body = _rewrite(GEN24, lambda ws: setattr(ws.cell(row=1, column=3), "value", "Early Voting"))
    with pytest.raises(SchemaDrift, match="unrecognised sheet layout"):
        tn.parse(body, 2024, LAST24)


def test_an_extra_column_is_drift():
    body = _rewrite(GEN24, lambda ws: setattr(ws.cell(row=1, column=6), "value", "Provisional"))
    with pytest.raises(SchemaDrift, match="unrecognised sheet layout"):
        tn.parse(body, 2024, LAST24)


def test_unknown_county_name_raises_drift():
    body = _rewrite(GEN24, lambda ws: setattr(ws.cell(row=15, column=1), "value", "Atlantis"))
    with pytest.raises(SchemaDrift, match="unrecognised county"):
        tn.parse(body, 2024, LAST24)


def test_non_numeric_count_raises_drift():
    body = _rewrite(GEN24, lambda ws: setattr(ws.cell(row=2, column=5), "value", "lots"))
    with pytest.raises(SchemaDrift, match="not a count"):
        tn.parse(body, 2024, LAST24)


def test_unparseable_date_raises_drift():
    body = _rewrite(GEN24, lambda ws: setattr(ws.cell(row=2, column=2), "value", "someday"))
    with pytest.raises(SchemaDrift, match="not a date"):
        tn.parse(body, 2024, LAST24)


def test_a_non_xlsx_body_is_a_source_error():
    with pytest.raises(SourceError):
        tn.parse(b"<!DOCTYPE html><html>nope</html>", 2026, date(2026, 10, 20))


# --------------------------------------------------------------------------
# Filenames: the election-date prefix is what separates primary from general
# --------------------------------------------------------------------------
def test_filename_templates_carry_the_election_date():
    assert tn.prefix(2024) == "20241105"
    assert tn.prefix(2026) == "20261103"
    names = tn.filenames(2024, date(2024, 10, 31))
    assert "20241105EarlyAbsentee1031.xlsx" in names
    assert "20241105EarlyAbsenteethrough1031.xlsx" in names
    # Tennessee's August state primary carries a different election date, so no
    # August file can ever match a November prefix.
    assert all(n.startswith("20241105") for n in names)


def test_candidates_walk_back_from_the_run_date(monkeypatch):
    """Tennessee posts on business days, so a Monday run must be able to reach
    back to Friday's file rather than reporting nothing."""
    monkeypatch.setattr(tn.TNScraper, "_index_urls", lambda self, cycle: [])
    names = [u.rsplit("/", 1)[-1]
             for u in tn.TNScraper()._candidates(2026, date(2026, 10, 20))]
    assert names[0] == "20261103EarlyAbsentee1020.xlsx", "newest first"
    assert "20261103EarlyAbsentee1010.xlsx" in names
    assert "20261103EarlyAbsentee1009.xlsx" not in names, "stops at LOOKBACK_DAYS"
    assert all(n.startswith("20261103") for n in names)


def test_candidates_never_look_past_election_day(monkeypatch):
    monkeypatch.setattr(tn.TNScraper, "_index_urls", lambda self, cycle: [])
    names = [u.rsplit("/", 1)[-1]
             for u in tn.TNScraper()._candidates(2026, date(2026, 12, 1))]
    assert names[0] == "20261103EarlyAbsentee1103.xlsx"


def test_index_discovery_only_accepts_this_cycles_prefix():
    index = (
        '<a href="/s3fs-public/document/20260806EarlyAbsentee0801.xlsx">Aug primary</a>'
        '<a href="/s3fs-public/document/20261103EarlyAbsentee1020.xlsx">general</a>'
        '<a href="/s3fs-public/document/20261103Turnout.xlsx">not the EV file</a>'
    )
    assert tn.discover(index, 2026) == [
        f"{tn.BASE}/20261103EarlyAbsentee1020.xlsx"
    ]
    assert tn.discover(index, 2024) == []
    assert tn.discover("", 2026) == []


# --------------------------------------------------------------------------
# Absence: an S3 403 means "no such file", not "we broke"
# --------------------------------------------------------------------------
def test_s3_403_is_read_as_absence():
    assert tn._absent(SourceError("TN: https://x returned HTTP 403"))
    assert not tn._absent(SourceError("TN: https://x returned HTTP 500"))
    assert not tn._absent(SourceError("TN: GET https://x failed: timeout"))


def test_no_workbook_anywhere_is_not_yet_published(monkeypatch):
    def missing(url, **kwargs):
        if url == tn.INDEX:
            raise Missing("TN: index 404")
        raise SourceError(f"TN: {url} returned HTTP 403")

    monkeypatch.setattr(tn, "get", missing)
    with pytest.raises(NotYetPublished, match="no early-voting workbook"):
        tn.TNScraper().fetch(2026, date(2026, 10, 20))


def test_a_real_outage_still_falls_through_the_ladder(monkeypatch):
    """A 500 is not absence. It must surface as SourceError so the runner tries
    the next tier instead of silently reporting 'nothing published yet'."""
    def broken(url, **kwargs):
        if url == tn.INDEX:
            raise Missing("TN: index 404")
        raise SourceError(f"TN: {url} returned HTTP 500")

    monkeypatch.setattr(tn, "get", broken)
    with pytest.raises(SourceError):
        tn.TNScraper().fetch(2026, date(2026, 10, 20))


def test_lookback_picks_up_an_earlier_days_file(monkeypatch):
    body = GEN24.read_bytes()
    wanted = "20261103EarlyAbsentee1016.xlsx"

    def get(url, **kwargs):
        if url == tn.INDEX:
            raise Missing("TN: index 404")
        if url.endswith(wanted):
            return body
        raise SourceError(f"TN: {url} returned HTTP 403")

    monkeypatch.setattr(tn, "get", get)
    result = tn.TNScraper().fetch(2026, date(2026, 10, 20))
    assert result.state_rows, "the older workbook was found and parsed"


def test_history_of_the_current_cycle_is_not_an_archive():
    with pytest.raises(NotYetPublished, match="not an archived cycle"):
        tn.TNScraper().fetch_history(2026)


def test_adapter_identity():
    scraper = tn.TNScraper()
    assert (scraper.state, scraper.name, scraper.tier) == ("TN", "tn-sos", 1)
