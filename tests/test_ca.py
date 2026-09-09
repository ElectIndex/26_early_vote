"""California: the SoS's VBM statistics report.

Five fixtures, all real, and between them every shape and vintage California has
published this report in:

* `vbm-statistics_2022-general_2022-10-27.xlsx` -- the DURING-SEASON snapshot,
  taken twelve days before the 2022 general. Its "VBM Press Version" sheet is
  copied cell for cell out of the 452 KB eight-sheet workbook the SoS served
  that day (Wayback capture `20221027003436`); only the seven sheets this
  adapter never reads are dropped, to keep the fixture small. This is the file
  that proves California publishes returns while the window is open.
* `vbm-statistics_2022-general_final.xlsm` -- the same URL as served live today,
  byte for byte, 77,507 bytes. It is the FINAL restatement and carries only the
  one sheet, which is why the parser finds its sheet by name.
* `vbm-statistics_2024-general_2024-10-31.xls` -- 129,536 bytes, legacy OLE2,
  the ONLY workbook capture of the 2024 general (Wayback `20241101232312`,
  `x-archive-orig-last-modified: Thu, 31 Oct 2024 16:15:00 GMT`). Its sheet is
  called `Ballot Return Statistics`, not `VBM Press Version`, and its header
  carries the in-person block -- including a SECOND column called `Sum`.
* `vbm-statistics_2024-general_2024-10-19.pdf` -- 124,309 bytes, the during-season
  PDF (Wayback `20241020170838`, last-modified 2024-10-19, printed "Ballot Return
  Statistics as of: Saturday, October 19, 2024"). Three of the five distinct 2024
  captures are PDFs, so this shape is most of the curve.
* `vbm-statistics_2022-general_2022-10-25.pdf` -- 168,846 bytes, the 2022 PDF of
  the same report, kept because it must be REFUSED. Its blank cells print as
  nothing, so Alameda's line carries ten numbers where its neighbours carry
  eleven and no reader can tell by position which channel went missing.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from ev.adapters import ca
from ev.adapters.base import NotYetPublished, SchemaDrift, SourceError

FIX = Path(__file__).parent / "fixtures" / "ca"
SEASON = (FIX / "vbm-statistics_2022-general_2022-10-27.xlsx").read_bytes()
FINAL = (FIX / "vbm-statistics_2022-general_final.xlsm").read_bytes()
SEASON_2024 = (FIX / "vbm-statistics_2024-general_2024-10-31.xls").read_bytes()
PDF_2024 = (FIX / "vbm-statistics_2024-general_2024-10-19.pdf").read_bytes()
PDF_2022 = (FIX / "vbm-statistics_2022-general_2022-10-25.pdf").read_bytes()

DAY = date(2022, 10, 27)
DAY_2024 = date(2024, 10, 31)
PDF_DAY_2024 = date(2024, 10, 19)


# --------------------------------------------------------------------------
# The during-season file, which the coverage survey said did not exist
# --------------------------------------------------------------------------
def test_the_during_season_snapshot_is_all_58_counties():
    result = ca.parse(SEASON, 2022, DAY)
    assert len(result.county_rows) == 58
    assert len({r.county_fips for r in result.county_rows}) == 58
    assert all(r.county_fips.startswith("06") for r in result.county_rows)


def test_the_statewide_row_is_californias_own_total():
    (row,) = ca.parse(SEASON, 2022, DAY).state_rows
    assert row.day == DAY
    assert row.ballots_total == 1_642_945
    assert row.mail_requested == 22_156_707


def test_the_counties_add_up_to_californias_own_total():
    result = ca.parse(SEASON, 2022, DAY)
    assert sum(r.ballots_total for r in result.county_rows) == 1_642_945


def test_los_angeles_is_keyed_by_fips_and_named_from_the_census():
    result = ca.parse(SEASON, 2022, DAY)
    la = next(r for r in result.county_rows if r.county_fips == "06037")
    assert la.county_name == "Los Angeles County"
    assert la.ballots_total == 310_707


def test_the_final_restatement_parses_from_its_single_sheet():
    """The live file has only the 'VBM Press Version' sheet; the during-season
    one has eight. The sheet is found by name for exactly that reason."""
    result = ca.parse(FINAL, 2022, date(2022, 11, 8))
    assert len(result.county_rows) == 58
    assert result.state_rows[0].ballots_total == 5_230_699


def test_returns_grow_between_the_two_snapshots():
    """This is what makes the source a tracker rather than a restatement."""
    early = ca.parse(SEASON, 2022, DAY).state_rows[0].ballots_total
    final = ca.parse(FINAL, 2022, date(2022, 11, 8)).state_rows[0].ballots_total
    assert early < final


# --------------------------------------------------------------------------
# What the counts mean
# --------------------------------------------------------------------------
def test_ballots_total_is_returns_not_the_accepted_subset():
    """Accepted lags returns by the county's signature-review queue, and a
    headline that moves when a county finishes checking is not turnout."""
    result = ca.parse(SEASON, 2022, DAY)
    la = next(r for r in result.county_rows if r.county_fips == "06037")
    # The sheet's own accepted figure for LA that day is lower than its Sum.
    assert la.ballots_total == 310_707
    assert ca.RETURNED == "sum" and ca.ACCEPTED == "total accepted vbm ballots"


def test_every_return_channel_is_mail_and_in_person_stays_blank():
    """Drop box, drop-off, vote-centre drop-off, post, fax and other are all a
    vote-by-mail ballot coming home. Ballots cast in person at a vote centre are
    not in this workbook at all, so that field is blank rather than 0."""
    for row in ca.parse(SEASON, 2022, DAY).county_rows:
        assert row.mail_returned == row.ballots_total
        assert row.inperson is None


def test_california_registers_by_party_but_this_workbook_does_not_split_it():
    result = ca.parse(SEASON, 2022, DAY)
    for row in result.state_rows + result.county_rows:
        assert (row.party_dem, row.party_rep, row.party_npa, row.party_oth) == (
            None, None, None, None
        )


# --------------------------------------------------------------------------
# The 2024 vintage: a renamed sheet, an in-person block, and TWO columns
# both called "Sum"
# --------------------------------------------------------------------------
def test_the_2024_workbook_is_all_58_counties():
    result = ca.parse(SEASON_2024, 2024, DAY_2024)
    assert len(result.county_rows) == 58
    assert len({r.county_fips for r in result.county_rows}) == 58


def test_the_second_sum_is_the_in_person_block_not_the_returns(monkeypatch):
    """⚠️ THE ONE THAT WOULD HAVE PUBLISHED CONFIDENT WRONG NUMBERS. The 2024
    header prints `Sum` twice. A last-occurrence map reads Los Angeles' returns
    as its in-person count -- 80,767 instead of 1,114,418, a factor of fourteen,
    in a column that looks entirely plausible on a chart."""
    la = next(r for r in ca.parse(SEASON_2024, 2024, DAY_2024).county_rows
              if r.county_fips == "06037")
    assert la.mail_returned == 1_114_418
    assert la.inperson == 80_767
    assert la.ballots_total == 1_195_185 == la.mail_returned + la.inperson


def test_index_binds_a_repeated_name_to_its_first_column():
    header = ["county", "sum", "regular ballots", "sum"]
    index = ca._index(header)
    assert index["sum"] == 1
    assert ca._inperson_sum(header, index) == 3


def test_a_vintage_with_no_in_person_block_has_no_in_person_column():
    header = ["county", "sum", "total accepted vbm ballots"]
    assert ca._inperson_sum(header, ca._index(header)) is None


def test_the_counties_add_up_to_the_2024_totals_in_both_transports():
    """The strongest check there is that the mapping is right, and it holds for
    the workbook and the PDF alike: California's own printed Total row, re-added
    from its own 58 counties, in all three published columns."""
    for body, day in ((SEASON_2024, DAY_2024), (PDF_2024, PDF_DAY_2024)):
        result = ca.parse(body, 2024, day)
        (state,) = result.state_rows
        assert sum(r.ballots_total for r in result.county_rows) == state.ballots_total
        assert sum(r.mail_returned for r in result.county_rows) == state.mail_returned
        assert sum(r.inperson for r in result.county_rows) == state.inperson


def test_the_2024_statewide_row_carries_both_methods():
    (row,) = ca.parse(SEASON_2024, 2024, DAY_2024).state_rows
    assert row.mail_requested == 22_835_999
    assert row.mail_returned == 5_832_308
    assert row.inperson == 177_868
    assert row.ballots_total == 6_010_176


def test_a_total_ballots_cast_that_does_not_add_up_is_drift():
    """The reconciliation that proves the column mapping rather than assuming
    it: California's own grand total, re-added from its own two Sums."""
    rows = ca._rows(SEASON_2024, 2024)
    header = [ca._norm(c) for c in rows[3]]
    at = ca._index(header)[ca.TOTAL_CAST]
    rows[4][at] = 12_345

    def fake(body, cycle):
        return rows

    import ev.adapters.ca as module
    real, module._rows = module._rows, fake
    try:
        with pytest.raises(SchemaDrift, match="TOTAL BALLOTS CAST"):
            ca.parse(SEASON_2024, 2024, DAY_2024)
    finally:
        module._rows = real


# --------------------------------------------------------------------------
# The PDF shape of the same table
# --------------------------------------------------------------------------
def test_the_2024_pdf_reads_as_the_same_rows_a_sheet_does():
    """Three of the five distinct 2024 captures are PDFs, so this shape is most
    of California's 2024 curve -- not a fallback."""
    result = ca.parse(PDF_2024, 2024, PDF_DAY_2024)
    assert len(result.county_rows) == 58
    la = next(r for r in result.county_rows if r.county_fips == "06037")
    assert (la.mail_returned, la.inperson, la.ballots_total) == (366_747, 193, 366_940)
    (state,) = result.state_rows
    assert (state.mail_returned, state.inperson, state.ballots_total) == (
        1_865_988, 3_316, 1_869_304
    )
    assert state.mail_requested == 22_729_631


def test_the_pdf_and_the_workbook_agree_about_what_a_row_means():
    """Same parser, same reconciliations, two transports. The PDF day is earlier
    than the workbook day, so every figure must be smaller and none may be a
    different KIND of figure."""
    pdf = ca.parse(PDF_2024, 2024, PDF_DAY_2024)
    book = ca.parse(SEASON_2024, 2024, DAY_2024)
    assert {r.county_fips for r in pdf.county_rows} == {
        r.county_fips for r in book.county_rows
    }
    assert pdf.state_rows[0].ballots_total < book.state_rows[0].ballots_total


def test_a_pdf_of_another_election_is_refused_by_name():
    """The Arizona bug, refused before it can happen: a report is only this
    cycle's if it says so on its own first line."""
    with pytest.raises(NotYetPublished, match="is the 2024 general, not 2022's"):
        ca.parse(PDF_2024, 2022, date(2022, 11, 8))


def test_the_2022_pdf_is_refused_rather_than_read_by_position():
    """⚠️ THE REASON THE PDF READER IS 2024-ONLY. 2022 leaves an unused return
    channel BLANK; a blank cell prints as nothing at all, so Alameda's line
    carries ten numbers where its neighbours carry eleven and nothing in the
    text says WHICH channel went missing. Reading it by position shifts every
    column left."""
    with pytest.raises(SchemaDrift, match=r"prints 10 fields, not 19"):
        ca.parse(PDF_2022, 2022, date(2022, 10, 25))


def test_the_pdf_field_count_is_a_hard_equality():
    assert ca.PDF_FIELDS == 19
    assert len(ca.PDF_HEADER) == 21
    assert ca.PDF_HEADER[0] == ca.COUNTY and ca.PDF_HEADER[-1] == ca.TOTAL_CAST


def test_a_data_line_is_found_by_county_type_not_by_county_name():
    """County names contain spaces; the three County Type values do not appear
    anywhere else on the page."""
    label, fields = ca._pdf_split("San Luis Obispo Polling Place 1 2 3")
    assert label == "San Luis Obispo" and fields == ["1", "2", "3"]
    assert ca._pdf_split("Accepted % of Voter-returned Ballots") == (None, [])


def test_the_statewide_line_is_read_with_or_without_its_label():
    """California rendered the same report two ways in 2024: in one the word
    "Total" sits with its numbers, in the other it is laid out with the header
    and the numbers print alone."""
    labelled, fields = ca._pdf_split("Total " + " ".join(["1"] * ca.PDF_FIELDS))
    assert labelled == "Total" and len(fields) == ca.PDF_FIELDS
    bare, fields = ca._pdf_split(" ".join(["1,000"] * ca.PDF_FIELDS))
    assert bare == "Total" and len(fields) == ca.PDF_FIELDS


# --------------------------------------------------------------------------
# The sheet has to prove it is the sheet we think it is
# --------------------------------------------------------------------------
def test_a_workbook_with_neither_known_sheet_name_is_drift():
    import io
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Something Else"
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="none of the sheets"):
        ca.parse(buf.getvalue(), 2022, DAY)


def test_both_of_californias_sheet_names_are_accepted():
    """⚠️ THE 2024 VINTAGE WAS UNREADABLE FOR WANT OF THIS. The module asserted
    that "VBM Press Version" was the only sheet common to every version of the
    file; 2024 shipped `Ballot Return Statistics` and nothing else."""
    assert ca.SHEETS == ("Ballot Return Statistics", "VBM Press Version")
    assert len(ca.parse(SEASON, 2022, DAY).county_rows) == 58
    assert len(ca.parse(SEASON_2024, 2024, DAY_2024).county_rows) == 58


def test_a_missing_column_is_drift():
    import io
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(SEASON))
    sheet = book[ca.SHEET]
    for cell in sheet[4]:
        if str(cell.value or "").strip().lower() == ca.RETURNED:
            cell.value = "Total Returned"
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="missing columns"):
        ca.parse(buf.getvalue(), 2022, DAY)


def test_a_county_whose_channels_do_not_add_to_its_sum_is_drift():
    import io
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(SEASON))
    sheet = book[ca.SHEET]
    header = [str(c.value or "").strip().lower() for c in sheet[4]]
    col = header.index("mail") + 1
    sheet.cell(row=5, column=col).value = 999_999
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="return channels add to"):
        ca.parse(buf.getvalue(), 2022, DAY)


def test_an_unrecognised_county_is_drift():
    import io
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(SEASON))
    book[ca.SHEET].cell(row=5, column=2).value = "Alamedaa"
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="unrecognised county names"):
        ca.parse(buf.getvalue(), 2022, DAY)


def test_a_short_workbook_publishes_counties_and_no_statewide_row():
    import io
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(SEASON))
    book[ca.SHEET].delete_rows(5)          # drop Alameda
    buf = io.BytesIO()
    book.save(buf)
    result = ca.parse(buf.getvalue(), 2022, DAY)
    assert result.state_rows == []
    assert len(result.county_rows) == 57


# --------------------------------------------------------------------------
# Transport: a missing file answers 403, not 404
# --------------------------------------------------------------------------
def test_an_s3_access_denied_body_is_absence_not_a_refusal():
    assert ca.looks_absent(
        b'<?xml version="1.0" encoding="UTF-8"?>\n'
        b"<Error><Code>AccessDenied</Code><Message>Access Denied</Message></Error>"
    )
    assert not ca.looks_absent(b"<!DOCTYPE html><title>Just a moment...</title>")


def test_a_403_without_an_s3_body_falls_through_rather_than_stopping(monkeypatch):
    class Response:
        status_code = 403
        content = b"<!DOCTYPE html><title>Just a moment...</title>"
        headers: dict[str, str] = {}

    monkeypatch.setattr(ca.SESSION, "get", lambda *a, **k: Response())
    with pytest.raises(SourceError, match="without an S3 AccessDenied body"):
        ca.download("https://example.invalid/x.xlsx", filename="x.xlsx")


def test_a_legacy_xls_body_is_downloaded_rather_than_refused(monkeypatch):
    """⚠️ IT USED TO BE REFUSED, and that is why California had no 2024.

    `download` rejected any body not beginning `PK`, so the `.xls` entry in
    EXTENSIONS, the `xlrd` dependency and `_rows`' OLE2 branch were all
    unreachable -- and 2024, which California published only as OLE2 and as PDF,
    could never be read whatever the parser could do with it."""
    class Response:
        status_code = 200
        content = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 8000
        headers = {"Last-Modified": "Fri, 01 Nov 2024 12:00:00 GMT"}

    monkeypatch.setattr(ca.SESSION, "get", lambda *a, **k: Response())
    body, written = ca.download("https://example.invalid/x.xls", filename="x.xls")
    assert body.startswith(ca._OLE2_MAGIC)
    assert written == date(2024, 11, 1)


def test_a_pdf_body_is_downloaded_rather_than_refused(monkeypatch):
    class Response:
        status_code = 200
        content = b"%PDF-1.7" + b"\0" * 8000
        headers = {"Last-Modified": "Sun, 20 Oct 2024 12:00:00 GMT"}

    monkeypatch.setattr(ca.SESSION, "get", lambda *a, **k: Response())
    body, _ = ca.download("https://example.invalid/x.pdf", filename="x.pdf")
    assert body.startswith(b"%PDF-")


def test_a_body_that_is_neither_a_workbook_nor_a_pdf_is_a_source_error(monkeypatch):
    class Response:
        status_code = 200
        content = b"<!doctype html>" + b" " * 8000
        headers: dict[str, str] = {}

    monkeypatch.setattr(ca.SESSION, "get", lambda *a, **k: Response())
    with pytest.raises(SourceError, match="not a workbook or a PDF"):
        ca.download("https://example.invalid/x.xlsx", filename="x.xlsx")


def test_the_as_of_date_is_the_objects_own_write_time():
    assert ca.last_modified("Mon, 16 Jun 2025 22:51:49 GMT") == date(2025, 6, 16)
    assert ca.last_modified("not a date") is None
    assert ca.last_modified(None) is None


# --------------------------------------------------------------------------
# Adapter wiring
# --------------------------------------------------------------------------
def test_adapter_identity():
    scraper = ca.CAScraper()
    assert (scraper.state, scraper.name, scraper.tier) == ("CA", "ca-sos", 1)


def test_the_url_is_built_from_the_cycle():
    assert ca.URL.format(cycle=2026, ext=".xlsx").endswith(
        "/statewide-elections/2026-general/vbm-statistics.xlsx"
    )


def test_a_stale_workbook_is_refused_rather_than_republished(monkeypatch):
    """The live 2022 workbook was rewritten on 2025-06-16. Publishing it as
    this season's snapshot would put three-year-old returns on today's curve."""
    monkeypatch.setattr(
        ca.CAScraper, "_load",
        lambda self, cycle, **k: (FINAL, date(2025, 6, 16)),
    )
    with pytest.raises(NotYetPublished, match="outside the window"):
        ca.CAScraper().fetch(2022, date(2022, 11, 8))


def test_a_workbook_with_no_last_modified_is_not_published(monkeypatch):
    monkeypatch.setattr(ca.CAScraper, "_load",
                        lambda self, cycle, **k: (SEASON, None))
    with pytest.raises(SourceError, match="without a Last-Modified header"):
        ca.CAScraper().fetch(2022, date(2022, 10, 27))


def test_fetch_stamps_the_row_with_the_objects_write_date(monkeypatch):
    monkeypatch.setattr(ca.CAScraper, "_load",
                        lambda self, cycle, **k: (SEASON, date(2022, 10, 27)))
    (row,) = ca.CAScraper().fetch(2022, date(2022, 10, 28)).state_rows
    assert row.day == date(2022, 10, 27)
    assert row.ballots_total == 1_642_945


def test_a_workbook_stamped_after_the_run_falls_through(monkeypatch):
    monkeypatch.setattr(ca.CAScraper, "_load",
                        lambda self, cycle, **k: (SEASON, date(2022, 10, 27)))
    with pytest.raises(SourceError, match="after the run date"):
        ca.CAScraper().fetch(2022, date(2022, 10, 26))


# --------------------------------------------------------------------------
# The archived series
# --------------------------------------------------------------------------
def test_a_cycle_that_is_not_over_has_nothing_to_backfill():
    """GUARD PARITY, the shape nd.py and wa.py both carry. `fetch` stamps its
    rows with the object's write date and refuses one after the run date; this
    path must refuse a live cycle BY NAME rather than harvesting today's
    position and filing it under a future Election Day."""
    with pytest.raises(NotYetPublished, match="is not over"):
        ca.CAScraper().fetch_history(2026)


def test_a_cycle_with_nothing_archived_is_refused_by_name(monkeypatch):
    monkeypatch.setattr(ca, "_archive_captures", lambda cycle: [])
    with pytest.raises(NotYetPublished, match="nothing readable is archived"):
        ca.CAScraper().fetch_history(2024)


def test_a_capture_without_the_original_last_modified_is_skipped(monkeypatch):
    """A crawl time is when we SAW the file, which is on or after when
    California wrote it. Dating a row by it puts the row on the wrong day, so a
    capture that kept no `x-archive-orig-last-modified` is dropped."""
    monkeypatch.setattr(ca, "_archive_captures",
                        lambda cycle: [("20241020170838", "https://x/vbm-statistics.pdf")])
    monkeypatch.setattr(ca, "download", lambda *a, **k: (PDF_2024, None))
    with pytest.raises(NotYetPublished, match="nothing readable is archived"):
        ca.CAScraper().fetch_history(2024)


def test_a_capture_stamped_outside_the_cycles_window_is_skipped(monkeypatch):
    """The live 2022 workbook carries a 2025-06-16 re-upload stamp. The archived
    copy of the same bytes carries its ORIGINAL one, which is the point -- but a
    capture that somehow does not is refused here exactly as `fetch` refuses it."""
    monkeypatch.setattr(ca, "_archive_captures",
                        lambda cycle: [("20250616000000", "https://x/vbm-statistics.xlsm")])
    monkeypatch.setattr(ca, "download", lambda *a, **k: (FINAL, date(2025, 6, 16)))
    with pytest.raises(NotYetPublished, match="nothing readable is archived"):
        ca.CAScraper().fetch_history(2022)


def test_the_archived_curve_is_one_day_per_capture_dated_by_last_modified(monkeypatch):
    captures = [("20241020170838", "https://x/vbm-statistics.pdf"),
                ("20241101232312", "https://x/vbm-statistics.xls")]
    bodies = {"20241020170838": (PDF_2024, date(2024, 10, 19)),
              "20241101232312": (SEASON_2024, date(2024, 10, 31))}
    monkeypatch.setattr(ca, "_archive_captures", lambda cycle: captures)
    monkeypatch.setattr(
        ca, "download",
        lambda url, **k: bodies[url.split("/web/", 1)[1].split("id_", 1)[0]],
    )
    result = ca.CAScraper().fetch_history(2024)
    assert sorted({r.day for r in result.county_rows}) == [
        date(2024, 10, 19), date(2024, 10, 31)
    ]
    assert len(result.county_rows) == 116
    assert len(result.state_rows) == 2
    # ...and the curve grows, which is what makes it a curve.
    early, late = sorted(result.state_rows, key=lambda r: r.day)
    assert early.ballots_total < late.ballots_total


def test_a_capture_of_a_layout_we_cannot_map_is_skipped_not_guessed(monkeypatch):
    """The 2022 PDFs. Skipping them costs two days of the 2022 curve; reading
    them by position would shift every column left on any county with a blank
    return channel."""
    monkeypatch.setattr(ca, "_archive_captures",
                        lambda cycle: [("20221027003352", "https://x/vbm-statistics.pdf")])
    monkeypatch.setattr(ca, "download", lambda *a, **k: (PDF_2022, date(2022, 10, 26)))
    with pytest.raises(NotYetPublished, match="nothing readable is archived"):
        ca.CAScraper().fetch_history(2022)


def test_the_cdx_sweep_dedupes_on_digest_and_keeps_each_captures_own_url(monkeypatch):
    """⚠️ THE QUERY STRING IS PART OF THE KEY. CloudFront ignored `?os=...` and
    served one object; the Archive keyed on the whole URL, so three of the five
    distinct 2024 versions live only under junk-query variants. Asking for the
    bare path at their timestamps gets a redirect to the nearest bare-path
    capture and collapses a five-day curve to two."""
    rows = [["timestamp", "original", "digest"],
            ["20241105224543", "https://c/vbm-statistics.pdf", "AAA"],
            ["20241106192325", "https://c/vbm-statistics.pdf?os=roku", "BBB"],
            ["20241107110213", "https://c/vbm-statistics.pdf?os=wtmb", "BBB"],
            ["20241020170838", "https://c/vbm-statistics.pdf", "CCC"]]

    class Response:
        status_code = 200
        text = __import__("json").dumps(rows)

    monkeypatch.setattr(ca.SESSION, "get", lambda *a, **k: Response())
    got = ca._archive_captures(2024)
    assert got == [("20241020170838", "https://c/vbm-statistics.pdf"),
                   ("20241105224543", "https://c/vbm-statistics.pdf"),
                   ("20241106192325", "https://c/vbm-statistics.pdf?os=roku")]


def test_an_empty_cdx_answer_is_absence_not_a_fault(monkeypatch):
    class Response:
        status_code = 200
        text = "  "

    monkeypatch.setattr(ca.SESSION, "get", lambda *a, **k: Response())
    assert ca._archive_captures(2024) == []


# --------------------------------------------------------------------------
# Nothing is blocking California, and that has to stay written down
# --------------------------------------------------------------------------
def test_california_is_empty_because_of_the_calendar_not_a_bot_wall():
    """⚠️ "CA has no county data in any cycle" reads like a block and is not one.

    Re-measured 2026-09-08 with plain `requests` and NO impersonation: every
    2026 path answers 403 with a 111-byte `<Code>AccessDenied</Code>` body from
    `Server: AmazonS3` -- S3's answer for a key that does not exist, because
    listing is denied -- and the 2022 workbook answers 200 with 77,507 real
    bytes. There is no WAF, no challenge and no fingerprint rule on this CDN,
    which is why this module has no impersonation path and should not grow one.
    """
    doc = ca.__doc__
    assert "NOTHING IS BLOCKING CALIFORNIA" in doc
    assert "AmazonS3" in doc
    # ...and the S3 absence body is still read as absence, not as a refusal.
    assert ca.looks_absent(b'<?xml version="1.0" encoding="UTF-8"?>'
                           b'<Error><Code>AccessDenied</Code>'
                           b'<Message>Access Denied</Message></Error>') is True


def test_a_403_that_is_not_s3_absence_is_a_refusal_not_absence():
    """The distinction the whole adapter turns on: S3 saying "no such key" STOPS
    the ladder, a WAF saying no falls THROUGH. A block page must never be read as
    California not having published."""
    from ev.adapters._net import WALL_MARKERS

    for marker in WALL_MARKERS:
        assert ca.looks_absent(marker + b" " * 400) is False


def test_the_archive_route_names_the_captures_it_reads():
    """The survey behind it must not have to be redone: every distinct capture,
    its served extension and the write date the Archive replayed for it."""
    doc = ca.CAScraper.fetch_history.__doc__
    for stamp in ("20241020170838", "20241101232312", "20241105142944",
                  "20241105224543", "20241106192325", "20221027003436",
                  "20221105004307", "20221108015555", "20221111024125"):
        assert stamp in doc
    assert "x-archive-orig-last-modified" in doc
