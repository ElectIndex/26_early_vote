"""Michigan: the SoS absentee-by-jurisdiction workbook.

Five fixtures, each the same four counties -- Alcona, Alger, Baraga, Keweenaw --
cut out of a real workbook and keeping its header and its statewide line
verbatim. Three of those four counties carry a "Less than 10" suppression in
2024 and none do in 2022, so the blank rule is exercised both ways.

They are five rather than one because Michigan has published this report under
FOUR different headers and written its statewide line in THREE different places,
twice changing inside a single cycle:

    2024-10-15  COUNTY | JURISDICTION | REQUESTS | ISSUED | RECEIVED
                statewide line: COUNTY == "TOTALS"; 2020 comparison sheet
    2024-10-22  same header, county names lose the " COUNTY" suffix
                statewide line: COUNTY blank, JURISDICTION == "TOTALS"
    2024-10-29  same again, statewide line JURISDICTION == "GRAND TOTAL"
    2022-10-17  <blank> | DLCOUNTYCODE | JURISDCODE | COUNTY | JURISDICTION |
                APPS RETURNED | BALLOTS SENT | BALLOTS RECEIVED
                statewide line: no label at all
    2022-10-24  <blank> | COUNTY | JURISDICTION | APPS RECEIVED | BALLOTS SENT |
                BALLOTS RECEIVED | ... | BALLOTS SPOILED | BALLOTS REJECTED

Every one of those is a way the old positional parser would have failed, and
three of them are how it did.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from ev.adapters import mi
from ev.adapters._net import Missing
from ev.adapters.base import NotYetPublished, SchemaDrift

FIXTURE = (Path(__file__).parent / "fixtures" / "mi"
           / "2024-10-15_General-Election-Data-by-Jurisdiction.xlsx")


@pytest.fixture(scope="module")
def body():
    return FIXTURE.read_bytes()


@pytest.fixture(scope="module")
def gen2024(body):
    return mi.parse(body, 2024)


# --------------------------------------------------------------------------
# THE BLANK RULE
# --------------------------------------------------------------------------
def test_no_party_field_is_ever_populated(gen2024):
    """Michigan has no party registration; a 0 would claim zero Democrats voted."""
    for row in gen2024.state_rows + gen2024.county_rows:
        for field in ("party_dem", "party_rep", "party_oth", "party_npa"):
            assert getattr(row, field) is None, (field, row)


# --------------------------------------------------------------------------
# Statewide comes from Michigan's own TOTALS row
# --------------------------------------------------------------------------
def test_statewide_is_michigans_totals_row_not_our_sum(gen2024):
    """The TOTALS row is computed from unsuppressed data, so it is exact where a
    sum of the published jurisdictions could only ever be a lower bound."""
    (state,) = gen2024.state_rows
    assert state.state == "MI"
    assert state.day == date(2024, 10, 15)
    assert state.ballots_total == 672585
    assert state.mail_returned == 672585
    assert state.mail_requested == 2112367  # ISSUED: ballots actually sent out
    # Far larger than the four counties in this fixture could ever sum to.
    assert state.ballots_total > sum(
        r.ballots_total or 0 for r in gen2024.county_rows) * 100


def test_michigan_reports_no_in_person_early_vote_column(gen2024):
    """The workbook has five columns and none of them is in-person early voting,
    so the field stays blank rather than being filled from the mail count."""
    (state,) = gen2024.state_rows
    assert state.inperson is None
    assert all(row.inperson is None for row in gen2024.county_rows)


# --------------------------------------------------------------------------
# "Less than 10": suppressed is unknown, not zero
# --------------------------------------------------------------------------
def test_county_with_a_suppressed_jurisdiction_is_blank_not_a_lower_bound(gen2024):
    """Alcona, Alger and Keweenaw each contain a jurisdiction Michigan masked as
    "Less than 10". Summing that as zero would understate Alger's issued count by
    up to 10%; blank correctly says Michigan did not tell us."""
    by_fips = {row.county_fips: row for row in gen2024.county_rows}
    for fips in ("26001", "26003", "26083"):  # Alcona, Alger, Keweenaw
        assert by_fips[fips].mail_returned is None
        assert by_fips[fips].ballots_total is None


def test_county_with_no_suppression_is_exact(gen2024):
    """Baraga's five townships are all above the floor, so it must be a number."""
    baraga = next(r for r in gen2024.county_rows if r.county_fips == "26013")
    assert baraga.county_name == "Baraga County"
    assert baraga.mail_returned == 602
    assert baraga.ballots_total == 602


def test_suppression_flag_is_recognised_not_parsed_as_a_number():
    assert mi._count("Less than 10") is mi.MASKED
    assert mi._count("less than 10") is mi.MASKED
    assert mi._count(1234) == 1234
    assert mi._count(None) is None
    with pytest.raises(SchemaDrift):
        mi._count("about a hundred")


def test_a_suppressed_component_blanks_only_the_measure_it_touches():
    assert mi._sum([1, 2, 3]) == 6
    assert mi._sum([1, mi.MASKED, 3]) is None
    assert mi._sum([None, None]) is None
    assert mi._sum([0, 0]) == 0  # a real reported zero survives


# --------------------------------------------------------------------------
# Jurisdictions roll up on the file's own COUNTY column
# --------------------------------------------------------------------------
def test_counties_are_fips_keyed_and_jurisdictions_do_not_leak(gen2024):
    """Michigan reports 1,520 cities and townships; none of them may reach the
    county table, and Novi exists as both a city and a township so a municipality
    crosswalk would have been wrong anyway."""
    assert {r.county_fips for r in gen2024.county_rows} == {
        "26001", "26003", "26013", "26083"}
    for row in gen2024.county_rows:
        assert len(row.county_fips) == 5 and row.county_fips.startswith("26")


def test_prior_cycle_sheet_is_selected_by_year(body):
    """The workbook carries a 2020 comparison sheet beside the live one; asking
    for a cycle must never silently return the other sheet's numbers."""
    old = mi.parse(body, 2020)
    assert {r.day for r in old.county_rows} == {date(2020, 10, 12)}
    alcona = next(r for r in old.county_rows if r.county_fips == "26001")
    assert alcona.mail_returned == 1340  # no suppression at all in the 2020 sheet


def test_sheet_title_supplies_year_and_as_of_date():
    assert mi.sheet_date("2024 (Oct 15)") == (2024, date(2024, 10, 15))
    assert mi.sheet_date("2020 (Oct 12)") == (2020, date(2020, 10, 12))
    with pytest.raises(SchemaDrift):
        mi.sheet_date("Absentee Data")


# --------------------------------------------------------------------------
# Drift and absence
# --------------------------------------------------------------------------
def test_renamed_column_raises_drift(tmp_path):
    import openpyxl

    book = openpyxl.load_workbook(FIXTURE)
    book["2024 (Oct 15)"].cell(row=1, column=5).value = "RETURNED"
    path = tmp_path / "drifted.xlsx"
    book.save(path)
    with pytest.raises(SchemaDrift, match="unrecognised column 'returned'"):
        mi.parse(path.read_bytes(), 2024)


def test_unknown_county_name_raises_drift(tmp_path):
    import openpyxl

    book = openpyxl.load_workbook(FIXTURE)
    book["2024 (Oct 15)"].cell(row=2, column=1).value = "ATLANTIS"
    path = tmp_path / "unknown.xlsx"
    book.save(path)
    with pytest.raises(SchemaDrift, match="unrecognised county"):
        mi.parse(path.read_bytes(), 2024)


def test_workbook_without_the_current_cycle_is_not_yet_published(body):
    """Before Michigan adds the 2026 sheet the file is still up, still valid, and
    still has nothing for us -- that is 'pending', not a failure."""
    with pytest.raises(NotYetPublished, match="no 2026 sheet"):
        mi.parse(body, 2026)


def test_missing_workbook_is_not_yet_published(monkeypatch):
    def missing(url, **kwargs):
        raise Missing(f"MI: {url} returned 404")

    monkeypatch.setattr(mi, "get", missing)
    with pytest.raises(NotYetPublished):
        mi.MIScraper().fetch(2026, date(2026, 9, 5))


def test_soft_404_html_page_is_not_yet_published(monkeypatch):
    monkeypatch.setattr(mi, "get", lambda url, **kw: b"<!DOCTYPE html><html>nope</html>")
    with pytest.raises(NotYetPublished):
        mi.MIScraper().fetch(2026, date(2026, 9, 5))


def test_future_dated_workbook_is_not_published(monkeypatch, body):
    monkeypatch.setattr(mi, "get", lambda url, **kw: body)
    with pytest.raises(NotYetPublished, match="after"):
        mi.MIScraper().fetch(2024, date(2024, 10, 1))


def test_adapter_identity():
    scraper = mi.MIScraper()
    assert (scraper.state, scraper.name, scraper.tier) == ("MI", "mi-sos", 1)


# ==========================================================================
# THE ARCHIVE: dated press-release copies of the same workbook
#
# Michigan overwrites `mi.URL` in place, but it attaches a DATED copy of the
# same workbook to each weekly early-vote press release. Those are the fixtures
# below, and between them they cover every header layout and every position the
# statewide line has been written in.
# ==========================================================================
FIXTURES = Path(__file__).parent / "fixtures" / "mi"

OCT22 = FIXTURES / "2024-10-22_10_22_2024-General-Election-Data-by-Jursidictionv3.xlsx"
OCT29 = FIXTURES / "2024-10-29_102924-General-Election-Data-by-Jursidiction.xlsx"
MID22 = FIXTURES / "2022-10-17_pr-nov-2022-ballot-stats-by-jurisdiction.xlsx"
LATE22 = FIXTURES / "2022-10-24_pr-nov-2022-av--ballot-stats.xlsx"


def test_totals_row_is_found_when_it_moves_into_the_jurisdiction_column():
    """2024-10-22 leaves COUNTY blank and writes "TOTALS" in JURISDICTION.

    The old loop skipped every row with a blank COUNTY, so it threw this away --
    and it is Michigan's own statewide figure, the one the press release quotes.
    """
    result = mi.parse(OCT22.read_bytes(), 2024)
    (state,) = result.state_rows
    assert state.day == date(2024, 10, 22)
    assert (state.mail_requested, state.mail_returned) == (2242984, 1147041)
    assert state.ballots_total == 1147041


def test_grand_total_is_the_same_statewide_line_under_another_name():
    """A week later Michigan calls it "GRAND TOTAL". Same row, same meaning."""
    result = mi.parse(OCT29.read_bytes(), 2024)
    (state,) = result.state_rows
    assert state.day == date(2024, 10, 29)
    assert (state.mail_requested, state.mail_returned) == (2328747, 1602831)


def test_2024_late_sheets_drop_the_county_suffix():
    """Oct 15 says "ALCONA COUNTY", Oct 22 and Oct 29 say "ALCONA". Both must
    reach the same FIPS -- which is the whole reason rows are keyed by FIPS."""
    for path in (OCT22, OCT29):
        rows = mi.parse(path.read_bytes(), 2024).county_rows
        assert {r.county_fips for r in rows} == {"26001", "26003", "26013", "26083"}
        assert all(r.county_name.endswith(" County") for r in rows)


def test_2022_eight_column_layout_is_read_by_name():
    """2022-10-17 leads with an unnamed index column and two internal code
    columns, and calls the measures APPS RETURNED / BALLOTS SENT / BALLOTS
    RECEIVED. Nothing about it is positionally the same as 2024."""
    result = mi.parse(MID22.read_bytes(), 2022)
    by_fips = {r.county_fips: r for r in result.county_rows}
    assert by_fips["26001"].day == date(2022, 10, 17)
    assert by_fips["26001"].mail_returned == 303
    assert by_fips["26083"].mail_returned == 70


def test_2022_padded_layout_ignores_spoiled_and_rejected_columns():
    """2022-10-24 pads out to twenty columns and ends with BALLOTS SPOILED and
    BALLOTS REJECTED. Both are named in IGNORED, so they neither break the
    header nor leak into a measure."""
    result = mi.parse(LATE22.read_bytes(), 2022)
    by_fips = {r.county_fips: r for r in result.county_rows}
    assert by_fips["26001"].mail_returned == 721
    assert by_fips["26013"].mail_returned == 504
    assert all(r.inperson is None for r in result.county_rows)


def test_an_unlabelled_trailing_total_is_dropped_rather_than_guessed():
    """The 2022 exports end with a row of statewide numbers and NO label at all.

    It is obviously the total -- but there is no label to normalise, and RULE 3
    is that a guess is worse than a gap. So the cycle gets county rows and no
    state row, and the state series comes from a weaker tier.
    """
    for path, cycle in ((MID22, 2022), (LATE22, 2022)):
        result = mi.parse(path.read_bytes(), cycle)
        assert result.county_rows
        assert result.state_rows == []


def test_layout_maps_every_real_header_michigan_has_published():
    assert mi.layout(("COUNTY", "JURISDICTION", "REQUESTS", "ISSUED", "RECEIVED")) == {
        "county": 0, "jurisdiction": 1, "requests": 2, "issued": 3, "received": 4}
    assert mi.layout((None, "DLCOUNTYCODE", "JURISDCODE", "COUNTY", "JURISDICTION",
                      "APPS RETURNED", "BALLOTS SENT", "BALLOTS RECEIVED")) == {
        "county": 3, "jurisdiction": 4, "requests": 5, "issued": 6, "received": 7}
    assert mi.layout((None, "COUNTY", "JURISDICTION", "APPS RECEIVED", "BALLOTS SENT",
                      "BALLOTS RECEIVED", None, "BALLOTS SPOILED",
                      "BALLOTS REJECTED")) == {
        "county": 1, "jurisdiction": 2, "requests": 3, "issued": 4, "received": 5}
    assert mi.layout(("COUNTY", "JURISDICTION", "APPS RETURNED", "BALLOTS SENT",
                      "BALLOTS RECEIVED")) == {
        "county": 0, "jurisdiction": 1, "requests": 2, "issued": 3, "received": 4}


def test_layout_refuses_a_name_it_does_not_know():
    with pytest.raises(SchemaDrift, match="unrecognised column 'ballots cured'"):
        mi.layout(("COUNTY", "JURISDICTION", "REQUESTS", "ISSUED", "RECEIVED",
                   "BALLOTS CURED"))


def test_layout_refuses_a_header_missing_a_measure():
    with pytest.raises(SchemaDrift, match=r"has no \['received'\]"):
        mi.layout(("COUNTY", "JURISDICTION", "REQUESTS", "ISSUED"))


def test_layout_refuses_a_measure_named_twice():
    with pytest.raises(SchemaDrift, match="appears twice"):
        mi.layout(("COUNTY", "JURISDICTION", "REQUESTS", "APPS RETURNED",
                   "ISSUED", "RECEIVED"))


def test_a_county_row_after_the_trailer_raises_drift(tmp_path):
    """If a row we treated as the sheet's trailing total is followed by more
    county data, it was not a trailer and everything after it is suspect."""
    import openpyxl

    book = openpyxl.load_workbook(OCT29)
    sheet = book["2024 (Oct 29)"]
    sheet.append(["ALCONA", "ALCONA TOWNSHIP", 1, 1, 1])
    path = tmp_path / "after_the_trailer.xlsx"
    book.save(path)
    with pytest.raises(SchemaDrift, match="after the sheet's trailing total"):
        mi.parse(path.read_bytes(), 2024)


def test_two_statewide_rows_on_one_sheet_raise_drift(tmp_path):
    import openpyxl

    book = openpyxl.load_workbook(FIXTURE)
    book["2024 (Oct 15)"].append(["TOTALS", None, 1, 2, 3])
    path = tmp_path / "twice.xlsx"
    book.save(path)
    with pytest.raises(SchemaDrift, match="two statewide rows"):
        mi.parse(path.read_bytes(), 2024)


# --------------------------------------------------------------------------
# fetch_history
# --------------------------------------------------------------------------
def test_fetch_history_refuses_a_cycle_with_no_pinned_archive():
    with pytest.raises(NotYetPublished, match="no archived 2020"):
        mi.MIScraper().fetch_history(2020)


def test_fetch_history_refuses_a_cycle_that_is_not_over(monkeypatch):
    """GUARD PARITY with nd.py and wa.py. ARCHIVE_PATHS is hand-pinned, so for a
    RUNNING cycle it is necessarily a partial list, and publishing a partial list
    as "the archive" would put a stub beside the live daily fetch."""
    monkeypatch.setitem(mi.ARCHIVE_PATHS, 2026, ("/nope.xlsx",))
    monkeypatch.setattr(mi, "election_date", lambda cycle: date(2999, 11, 3))
    with pytest.raises(NotYetPublished, match="2026 general has not happened yet"):
        mi.MIScraper().fetch_history(2026)


def _archive_stub(monkeypatch, bodies, captures_per_path=1):
    """Serve `bodies` (path -> bytes) instead of the Wayback Machine."""
    served: list[str] = []

    def captures(self, path, cycle):
        return [f"https://web.archive.org/web/2025090400000{n}id_/x{path}"
                for n in range(captures_per_path)]

    def get(url, **kwargs):
        served.append(url)
        for path, body in bodies.items():
            if url.endswith(path):
                if body is None:
                    raise Missing(f"MI: {url} returned 404")
                return body
        raise Missing(f"MI: {url} returned 404")

    monkeypatch.setattr(mi.MIScraper, "_captures", captures)
    monkeypatch.setattr(mi._net, "get", get)
    return served


def test_fetch_history_walks_the_pinned_archive(monkeypatch):
    paths = mi.ARCHIVE_PATHS[2024]
    _archive_stub(monkeypatch, {
        paths[0]: FIXTURE.read_bytes(),
        paths[1]: OCT22.read_bytes(),
        paths[2]: OCT29.read_bytes(),
    })
    result = mi.MIScraper().fetch_history(2024)
    assert [r.day for r in result.state_rows] == [
        date(2024, 10, 15), date(2024, 10, 22), date(2024, 10, 29)]
    assert {r.day for r in result.county_rows} == {
        date(2024, 10, 15), date(2024, 10, 22), date(2024, 10, 29)}
    # Four counties per fixture day, every one FIPS-keyed.
    assert len(result.county_rows) == 12
    assert all(len(r.county_fips) == 5 for r in result.county_rows)


def test_fetch_history_retries_a_capture_that_404s_on_playback(monkeypatch):
    """A capture the CDX index calls 200 can still 404 on playback -- the real
    2022-10-24 workbook does exactly that on its 2022 capture and serves fine on
    its 2025 one. Taking only the newest capture silently cost a whole day."""
    paths = mi.ARCHIVE_PATHS[2024]
    bodies = {p: None for p in paths}
    bodies[paths[2]] = OCT29.read_bytes()
    served = _archive_stub(monkeypatch, bodies, captures_per_path=3)
    result = mi.MIScraper().fetch_history(2024)
    assert [r.day for r in result.state_rows] == [date(2024, 10, 29)]
    # All three captures tried for each of the two paths that never answer, and
    # exactly one for the path whose first capture works -- the walk stops at the
    # first readable capture rather than fetching the rest.
    assert len(served) == 3 + 3 + 1


def test_fetch_history_with_nothing_readable_is_not_yet_published(monkeypatch):
    _archive_stub(monkeypatch, {p: None for p in mi.ARCHIVE_PATHS[2024]})
    with pytest.raises(NotYetPublished, match="nothing usable archived"):
        mi.MIScraper().fetch_history(2024)


def test_archive_paths_only_hold_finished_cycles():
    """The list is hand-pinned; a running cycle in it would be a partial archive
    competing with the live daily fetch. The guard refuses it, and this keeps
    anyone from adding one and only finding out from the guard."""
    from ev.calendar import election_date

    for cycle in mi.ARCHIVE_PATHS:
        assert election_date(cycle) < date.today(), cycle
