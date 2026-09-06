"""Kentucky: the SBE's dated daily absentee workbook.

Two fixtures, both truncated from the real files with rows 1-10 -- the SBE's
prose caveats and the real header -- kept verbatim:

* `Absentee_Public_110124.xlsx`, four days before the 2024 general, with every
  measure populated.
* `Absentee_Public_100824.xlsx`, four weeks out, when in-person voting had not
  opened. Its county in-person cells are BLANK while its statewide ones are a
  literal 0 -- the exact distinction THE BLANK RULE exists for, in one file.

Expected numbers below were added up by hand from the fixture's own cells.
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

from ev.adapters import ky
from ev.adapters._net import Missing
from ev.adapters.base import NotYetPublished, SchemaDrift

FIXTURES = Path(__file__).parent / "fixtures" / "ky"
LATE = FIXTURES / "Absentee_Public_110124.xlsx"
EARLY = FIXTURES / "Absentee_Public_100824.xlsx"

NOV1 = date(2024, 11, 1)
OCT8 = date(2024, 10, 8)


@pytest.fixture(scope="module")
def late():
    return ky.parse(LATE.read_bytes(), 2024, NOV1)


@pytest.fixture(scope="module")
def early():
    return ky.parse(EARLY.read_bytes(), 2024, OCT8)


# --------------------------------------------------------------------------
# THE BLANK RULE, in both directions -- the point of the second fixture
# --------------------------------------------------------------------------
def test_unreported_party_buckets_are_blank_never_zero(late, early):
    """Kentucky splits out Democrats and Republicans and nothing else. The
    residual mixes unaffiliated voters with third parties, so neither NPA nor
    OTH is derivable -- and a 0 would claim no unaffiliated Kentuckian voted."""
    for result in (late, early):
        for row in result.state_rows + result.county_rows:
            assert row.party_npa is None, row
            assert row.party_oth is None, row


def test_reported_party_buckets_are_real_counts(late):
    (state,) = late.state_rows
    assert state.party_dem == 50195 + 10291 + 89782     # mail + excused + no-excuse
    assert state.party_rep == 40153 + 17867 + 121081
    assert state.party_dem + state.party_rep < state.ballots_total


def test_blank_county_cells_stay_blank_but_a_reported_zero_stays_zero(early):
    """On 2024-10-08 in-person voting had not opened: Kentucky left the county
    cells empty and wrote 0 statewide. Those are different claims and the file
    makes both, so we must publish both."""
    (state,) = early.state_rows
    assert state.inperson == 0, "Kentucky affirmatively reported zero statewide"
    adair = {r.county_fips: r for r in early.county_rows}["21001"]
    assert adair.inperson is None, "an empty county cell is not a zero"
    assert adair.ballots_total == 77
    assert adair.mail_returned == 77 + 0


# --------------------------------------------------------------------------
# The statewide row is Kentucky's own, not our sum
# --------------------------------------------------------------------------
def test_statewide_row_is_kentuckys_own_number(late):
    (state,) = late.state_rows
    assert state.state == "KY"
    assert state.day == NOV1
    assert state.ballots_total == 355909              # verbatim from TOTALS
    assert state.mail_requested == 130464 + 4646      # domestic sent + FPCA
    assert state.mail_returned == 97842 + 2365        # domestic + FPCA returned
    assert state.inperson == 30006 + 225696           # excused + no-excuse
    assert state.mail_returned + state.inperson == state.ballots_total


def test_statewide_is_not_our_sum_of_the_counties(late):
    """The fixture keeps ten counties out of 120, so a parser that summed them
    would report 95,610 rather than Kentucky's 355,909."""
    (state,) = late.state_rows
    assert sum(r.ballots_total for r in late.county_rows) == 95610
    assert state.ballots_total == 355909


def test_totals_is_not_emitted_as_a_county(late):
    assert "TOTALS" not in {r.county_name for r in late.county_rows}
    assert all(r.county_fips != "21000" for r in late.county_rows)


def test_fpca_is_added_not_dropped(late):
    """Military and overseas ballots are their own block; Kentucky's own total
    adds them, which is what proves 'All Ballots RETURNED' is domestic-only."""
    (state,) = late.state_rows
    assert state.mail_returned > 97842


# --------------------------------------------------------------------------
# County rows are keyed by FIPS, never by name
# --------------------------------------------------------------------------
def test_county_rows_are_fips_keyed(late):
    by_fips = {r.county_fips: r for r in late.county_rows}
    assert by_fips["21111"].county_name == "Jefferson County"
    assert by_fips["21067"].county_name == "Fayette County"
    adair = by_fips["21001"]
    assert adair.ballots_total == 1334
    assert adair.mail_returned == 317 + 3
    assert adair.inperson == 88 + 926
    assert adair.mail_returned + adair.inperson == adair.ballots_total
    assert adair.party_dem == 81 + 17 + 156
    assert adair.party_rep == 221 + 64 + 714


def test_the_mc_and_la_counties_resolve(late):
    """"MCCRACKEN"/"MCCREARY"/"LARUE" are the names a case-sensitive or
    space-sensitive join gets wrong."""
    names = {r.county_fips: r.county_name for r in late.county_rows}
    assert names["21145"] == "McCracken County"
    assert names["21147"] == "McCreary County"
    assert names["21123"] == "Larue County"


def test_all_county_fips_are_five_digit_kentucky(late, early):
    for result in (late, early):
        for row in result.county_rows:
            assert len(row.county_fips) == 5 and row.county_fips.startswith("21")


def test_every_county_row_carries_the_files_own_date(late):
    assert {r.day for r in late.county_rows} == {NOV1}


# --------------------------------------------------------------------------
# The header is not on row 1, and "County" appears four times
# --------------------------------------------------------------------------
def test_header_is_found_below_the_prose_caveats():
    rows = list(openpyxl.load_workbook(LATE, data_only=True).active.iter_rows(values_only=True))
    at, index = ky._find_header(rows)
    assert at == 9, "row 10, one-based -- under nine rows of SBE prose"
    assert index[ky.COUNTY] == 0, "the first of four County columns"
    assert index[ky.NOEXCUSE] == 18
    assert index[ky.TOTAL] == 26


def test_missing_header_row_raises_drift():
    book = openpyxl.load_workbook(LATE)
    book.active.cell(row=10, column=1).value = "Jurisdiction"
    import io
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="no 'County' header row"):
        ky.parse(buf.getvalue(), 2024, NOV1)


def test_renamed_measure_column_raises_drift():
    book = openpyxl.load_workbook(LATE)
    book.active.cell(row=10, column=19).value = "Walk-in voting"
    import io
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="missing columns"):
        ky.parse(buf.getvalue(), 2024, NOV1)


def test_duplicated_measure_header_raises_drift():
    """If the four side-by-side blocks were rearranged we could not tell which
    one we were reading, so a repeated measure name must fail loudly."""
    book = openpyxl.load_workbook(LATE)
    book.active.cell(row=10, column=19).value = "Excused In-person"
    import io
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="duplicate header"):
        ky.parse(buf.getvalue(), 2024, NOV1)


def test_unknown_county_name_raises_drift():
    book = openpyxl.load_workbook(LATE)
    book.active.cell(row=11, column=1).value = "Atlantis"
    import io
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="unrecognised county"):
        ky.parse(buf.getvalue(), 2024, NOV1)


def test_non_numeric_cell_raises_drift():
    book = openpyxl.load_workbook(LATE)
    book.active.cell(row=11, column=8).value = "lots"
    import io
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(SchemaDrift, match="not a count"):
        ky.parse(buf.getvalue(), 2024, NOV1)


# --------------------------------------------------------------------------
# Filenames, the lookback, and the primary/general trap
# --------------------------------------------------------------------------
def test_filename_stamp():
    assert ky.stamp(date(2024, 11, 1)) == "110124"
    assert ky.stamp(date(2026, 5, 18)) == "051826"
    assert ky.stamp(date(2026, 11, 3)) == "110326"


def test_a_run_before_the_general_window_is_not_yet_published():
    """Kentucky posts the SAME filename pattern for its May primary, and those
    files are still on the server. Reading one in a 2026 run would publish
    primary turnout as general turnout."""
    scraper = ky.KYScraper()
    with pytest.raises(NotYetPublished, match="does not open"):
        scraper.fetch(2026, date(2026, 5, 18))
    with pytest.raises(NotYetPublished, match="does not open"):
        scraper.fetch(2026, date(2026, 7, 1))


def test_no_file_in_the_lookback_window_is_not_yet_published(monkeypatch):
    """This is the path that runs every day for weeks. It must STOP the ladder,
    not fall through to a source that would invent a zero."""
    def missing(url, **kwargs):
        raise Missing(f"KY: {url} returned 404")

    monkeypatch.setattr(ky, "get", missing)
    with pytest.raises(NotYetPublished, match="no absentee workbook"):
        ky.KYScraper().fetch(2026, date(2026, 10, 20))


def test_lookback_picks_up_a_file_from_an_earlier_day(monkeypatch):
    """Kentucky skips days -- 2024-10-15 and 2024-10-22 have no file while
    2024-10-08 does -- so a run must not report 'nothing yet' when a perfectly
    good snapshot from earlier in the week is sitting on the server."""
    body = LATE.read_bytes()
    wanted = f"Absentee_Public_{ky.stamp(date(2026, 10, 17))}.xlsx"
    asked: list[str] = []

    def get(url, **kwargs):
        asked.append(url)
        if url.endswith(wanted):
            return body
        raise Missing(f"KY: {url} returned 404")

    monkeypatch.setattr(ky, "get", get)
    result = ky.KYScraper().fetch(2026, date(2026, 10, 20))
    assert {r.day for r in result.state_rows} == {date(2026, 10, 17)}
    assert result.state_rows[0].cycle == 2026
    assert len(asked) == 4, "asked newest-first and stopped at the first hit"


def test_lookback_never_reaches_back_past_the_window(monkeypatch):
    """The walk back must stop at the window edge rather than stepping into the
    primary's files."""
    def missing(url, **kwargs):
        raise Missing(f"KY: {url} returned 404")

    monkeypatch.setattr(ky, "get", missing)
    start = date(2026, 11, 3) - timedelta(days=ky.WINDOW_DAYS)
    with pytest.raises(NotYetPublished):
        ky.KYScraper().fetch(2026, start + timedelta(days=2))


def test_html_error_page_counts_as_no_file(monkeypatch):
    monkeypatch.setattr(ky, "get", lambda url, **kw: b"<!DOCTYPE html><html>404</html>")
    with pytest.raises(NotYetPublished):
        ky.KYScraper().fetch(2026, date(2026, 10, 20))


def test_history_of_the_current_cycle_is_not_an_archive():
    with pytest.raises(NotYetPublished, match="not an archived cycle"):
        ky.KYScraper().fetch_history(2026)


def test_history_collects_every_day_that_exists(monkeypatch):
    """Kentucky's dated filenames make the archive a real daily series, unlike a
    state that overwrites one report in place."""
    posted = {
        ky.stamp(date(2024, 10, 8)): EARLY.read_bytes(),
        ky.stamp(date(2024, 11, 1)): LATE.read_bytes(),
    }

    def get(url, **kwargs):
        for mark, body in posted.items():
            if url.endswith(f"Absentee_Public_{mark}.xlsx"):
                return body
        raise Missing(f"KY: {url} returned 404")

    monkeypatch.setattr(ky, "get", get)
    result = ky.KYScraper().fetch_history(2024)
    assert sorted(r.day for r in result.state_rows) == [OCT8, NOV1]
    assert [r.ballots_total for r in sorted(result.state_rows, key=lambda r: r.day)] == [
        15097, 355909
    ]


def test_adapter_identity():
    scraper = ky.KYScraper()
    assert (scraper.state, scraper.name, scraper.tier) == ("KY", "ky-sbe", 1)
