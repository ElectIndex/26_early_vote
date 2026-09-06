"""Ohio: the SoS county absentee workbook.

Two fixtures, deliberately: the 2022 general report has 11 columns and the 2024
general report has 17, because Ohio added a dropbox/personal-delivery/by-mail
breakdown between cycles. Parsing both with the same code is the whole point of
matching columns by header name, so both are locked in here.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from ev.adapters import oh
from ev.adapters._net import Missing
from ev.adapters.base import NotYetPublished, SchemaDrift

FIXTURES = Path(__file__).parent / "fixtures" / "oh"


def _parse(name: str, cycle: int):
    return oh.parse((FIXTURES / name).read_bytes(), cycle)


@pytest.fixture(scope="module")
def gen2024():
    return _parse("2024gen_absentee_report_web.xlsx", 2024)


@pytest.fixture(scope="module")
def gen2022():
    return _parse("2022gen_absentee_report_web.xlsx", 2022)


# --------------------------------------------------------------------------
# THE BLANK RULE -- the one that matters most for this batch
# --------------------------------------------------------------------------
def test_no_party_field_is_ever_populated(gen2024, gen2022):
    """Ohio has no party registration. A 0 here would render as 'zero Democrats
    have voted' on a page whose whole subject is who is voting early."""
    for result in (gen2024, gen2022):
        for row in result.state_rows + result.county_rows:
            for field in ("party_dem", "party_rep", "party_oth", "party_npa"):
                assert getattr(row, field) is None, (field, row)


# --------------------------------------------------------------------------
# The statewide row is Ohio's own, not our sum
# --------------------------------------------------------------------------
def test_statewide_row_is_ohios_own_number(gen2024):
    (state,) = gen2024.state_rows
    assert state.state == "OH"
    assert state.day == date(2024, 11, 5)
    assert state.ballots_total == 2620750          # verbatim from the Statewide row
    assert state.mail_requested == 1131698 + 21870  # domestic + UOCAVA transmitted
    assert state.mail_returned == 1066815 + 17681   # domestic + UOCAVA by mail
    assert state.inperson == 1536219 + 35           # domestic + UOCAVA in person


def test_uocava_is_added_not_dropped(gen2024):
    """Military and overseas ballots live in their own column family; reading only
    the 'Domestic ...' columns silently understates every returns number."""
    (state,) = gen2024.state_rows
    assert state.mail_returned > 1066815


def test_statewide_is_not_emitted_as_a_county(gen2024):
    assert all(row.county_fips != "39000" for row in gen2024.county_rows)
    assert "Statewide" not in {row.county_name for row in gen2024.county_rows}


# --------------------------------------------------------------------------
# County rows are keyed by FIPS, never by name
# --------------------------------------------------------------------------
def test_county_rows_are_fips_keyed(gen2024):
    by_fips = {row.county_fips: row for row in gen2024.county_rows}
    assert by_fips["39035"].county_name == "Cuyahoga County"
    assert by_fips["39049"].county_name == "Franklin County"
    adams = by_fips["39001"]
    assert adams.ballots_total == 6966
    assert adams.mail_returned == 1681 + 20
    assert adams.inperson == 5265 + 0


def test_all_county_fips_are_five_digit_ohio(gen2024):
    for row in gen2024.county_rows:
        assert len(row.county_fips) == 5 and row.county_fips.startswith("39")


# --------------------------------------------------------------------------
# The 2022 layout has six fewer columns and must still parse
# --------------------------------------------------------------------------
def test_2022_layout_parses_by_name_not_position(gen2022):
    (state,) = gen2022.state_rows
    assert state.day == date(2022, 11, 8)
    assert state.ballots_total == 1473983
    assert state.mail_returned == 918333 + 5546
    assert state.inperson == 550092 + 12
    assert {r.county_fips for r in gen2022.county_rows} >= {"39001", "39035", "39049"}


def test_sheet_title_supplies_the_as_of_date():
    assert oh.sheet_date("11.5.2024 Absentee Report", 2024) == date(2024, 11, 5)
    assert oh.sheet_date("11.8.22 Absentee Report", 2022) == date(2022, 11, 8)


def test_sheet_title_from_the_wrong_cycle_is_drift():
    """A stale workbook left on the server would otherwise publish 2024 counts
    under the 2026 cycle."""
    with pytest.raises(SchemaDrift):
        oh.sheet_date("11.5.2024 Absentee Report", 2026)


# --------------------------------------------------------------------------
# Drift and absence
# --------------------------------------------------------------------------
def test_renamed_column_raises_drift(tmp_path):
    """Ohio renaming a column must fail loudly, not silently map to something."""
    import openpyxl

    book = openpyxl.load_workbook(FIXTURES / "2024gen_absentee_report_web.xlsx")
    sheet = book.worksheets[0]
    sheet.cell(row=1, column=4).value = "Ballots we mailed out"
    path = tmp_path / "drifted.xlsx"
    book.save(path)
    with pytest.raises(SchemaDrift, match="missing columns"):
        oh.parse(path.read_bytes(), 2024)


def test_unknown_county_name_raises_drift(tmp_path):
    """A name we cannot place would otherwise vanish off the county map."""
    import openpyxl

    book = openpyxl.load_workbook(FIXTURES / "2024gen_absentee_report_web.xlsx")
    book.worksheets[0].cell(row=3, column=1).value = "Atlantis"
    path = tmp_path / "unknown.xlsx"
    book.save(path)
    with pytest.raises(SchemaDrift, match="unrecognised county"):
        oh.parse(path.read_bytes(), 2024)


def test_missing_2026_report_is_not_yet_published(monkeypatch):
    """This is the path that runs every day for weeks before Ohio posts anything.
    It must STOP the ladder, not fall through to a source that invents a zero."""
    def missing(url, **kwargs):
        raise Missing(f"OH: {url} returned 404")

    monkeypatch.setattr(oh, "get", missing)
    with pytest.raises(NotYetPublished):
        oh.OHScraper().fetch(2026, date(2026, 9, 5))


def test_soft_404_html_page_is_not_yet_published(monkeypatch):
    """The SoS CMS answers an unposted report with a styled 200 HTML page."""
    monkeypatch.setattr(oh, "get", lambda url, **kw: b"<!DOCTYPE html><html>nope</html>")
    with pytest.raises(NotYetPublished):
        oh.OHScraper().fetch(2026, date(2026, 9, 5))


def test_future_dated_report_is_not_published(monkeypatch):
    body = (FIXTURES / "2024gen_absentee_report_web.xlsx").read_bytes()
    monkeypatch.setattr(oh, "get", lambda url, **kw: body)
    with pytest.raises(NotYetPublished, match="after"):
        oh.OHScraper().fetch(2024, date(2024, 10, 1))


def test_adapter_identity():
    scraper = oh.OHScraper()
    assert (scraper.state, scraper.name, scraper.tier) == ("OH", "oh-sos", 1)
