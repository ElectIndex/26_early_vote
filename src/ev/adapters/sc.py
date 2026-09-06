"""South Carolina: the State Election Commission's two daily workbooks.

South Carolina publishes its early vote as TWO files, both posted to the SEC's
WordPress media library most mornings of the season, and this adapter reads both:

* **`…Absentee-Statistics-by-Race…xlsx`** -- absentee-by-mail, one row per county,
  with ballots issued and ballots returned, split across nine per-race worksheets.
  It starts about six weeks out and runs to the end.
* **`…GE-Early-Voting-Statistics…xlsx`** -- in-person early voting, one row per
  county (with its polling places nested underneath) and ONE COLUMN PER DAY. Like
  North Carolina's and Maryland's files, one download reconstructs the whole
  in-person curve, so a missed run cannot lose a day.

In 2024 the two together are 1,185,345 early ballots: 1,106,201 in person and
79,144 by mail. Neither file alone is South Carolina's early vote.

**South Carolina does NOT register voters by party**, so all four party fields are
None on every row -- never 0. See THE BLANK RULE in schema.py.

Four things drive the code below.

* **Filenames drift hard, so the WordPress media API is the index.**
  `scvotes.gov/wp-json/wp/v2/media?search=…` returns every posted file with its
  upload timestamp, which is both the discovery mechanism and the snapshot's
  as-of date. That sidesteps a naming convention that changed from
  `Absentee-Stats-2022-10-21-noon-GE.xlsx` in 2022 to
  `2024-10-29-Absentee-Statistics-by-Race-for-GE.xlsx` in 2024, is not
  zero-padded consistently, and carries WordPress `-N` dedupe suffixes.

* **The general election, never a primary -- and the workbook itself is the
  guard.** Every workbook names its own election above the header: 2024 says
  `22114 Statewide General Election / Election Date: 11/5/2024`, and the file
  posted in August 2026 says `08/11/2026  22760  US Senate Special Republican
  Primary (RUNOFF)`. A candidate file is accepted only if that text carries the
  cycle's Election Day and does not say "primary". The filename is used only to
  ORDER candidates, never to decide.

* **Two report layouts, and the columns are found structurally.** The 2024
  report (`RP0007`) keys counties in column B; the 2026 one (`RP00010`) keys them
  in column F and adds a `Political Party:` filter line. So the header row is
  located by its "County" cell, and the measure columns are read off the two
  banner rows above it -- group (`Total` / `Non-UOCAVA` / `UOCAVA`) then
  sub-group (`Applications` / `Ballots`) -- rather than by position. Anything
  that does not line up raises SchemaDrift.

* **Race is published only while absentee IS the early vote.** The race
  breakdown covers absentee ballots only. Before in-person early voting opens
  that is the entire early vote and the breakdown is exact; after it opens the
  same numbers describe 7% of South Carolina's early voters (79,144 of
  1,185,345 in 2024) and would render as if they described all of them. So the
  race dimension is emitted only for days before the statutory in-person window
  opens -- the same call Maryland makes about its age bands, for the same reason.
"""

from __future__ import annotations

import io
import json
import logging
import re
from datetime import date, datetime, timedelta

import openpyxl

from ..calendar import election_date
from ..normalize import race as _race
from ..schema import TIER_SCRAPER, CountyDay, DemoDay, StateDay
from . import _fips
from ._net import Missing, get, looks_like_xlsx
from .base import Adapter, FetchResult, NotYetPublished, SchemaDrift, SourceError

log = logging.getLogger(__name__)

#: VERIFIED on this network: the media endpoint returns JSON listing every
#: absentee and early-voting workbook back to 2022, and each `source_url` it
#: names downloads as a real xlsx (200, e.g. the 74,282-byte
#: 2024-10-29 absentee report and the 28,152-byte 2024-10-31 early-voting one).
BASE = "https://scvotes.gov"
MEDIA = f"{BASE}/wp-json/wp/v2/media"

#: WordPress full-text search terms, run in order and unioned. Both naming eras
#: are covered: "Absentee-Stats-…" (2022) and "Absentee-Statistics-…" (2024/26).
SEARCH_TERMS = ("absentee statistics", "absentee stats", "early voting statistics")

#: How many candidate workbooks a single run will download before giving up.
#: The election-identity check inside the file is what actually decides, so this
#: only bounds the cost of a season where the filename hint stops working.
MAX_PROBE = 4

#: South Carolina's in-person early-voting window is statutory: the second Monday
#: before Election Day through the Saturday before it. Verified against all three
#: cycles -- 2022-10-24..11-05, 2024-10-21..11-02, 2026-10-19..10-31 -- and the
#: 2024 file's own date columns start on 10/21, exactly EV_OPEN_OFFSET days out.
EV_OPEN_OFFSET = 15
EV_CLOSE_OFFSET = 3

#: How many days BEHIND the absentee snapshot the in-person file may be. It may be
#: arbitrarily far ahead -- only its columns dated on or before the snapshot day
#: are ever summed, which is what makes a backfill work off one download -- but a
#: file generated days earlier is missing days of voting, and folding it in would
#: publish a total short by however many.
EV_MAX_SKEW_DAYS = 2

COUNTY = "county"
TOTAL_LABEL = "total"
GROUP_TOTAL = "total"
SUB_BALLOTS = "ballots"
BALLOT_MEASURES = ("issued", "returned before deadline",
                   "returned after deadline", "total returned")

#: Race labels `normalize.race()` does not carry. Both are real labels on the
#: SEC's own worksheets, so raising drift on them would leave South Carolina
#: permanently broken; consulted ONLY after normalize returns None, and anything
#: in neither still raises. Reported upstream for normalize.py.
EXTRA_RACE = {"black/aa": "black", "multiple": "other"}

HEADER_SEARCH_ROWS = 30

#: A real absentee report is 50-odd rows. Anything shorter than this with no
#: "County" header at all is the SEC's empty shell -- it posted one on
#: 2024-09-30, banner and all, with "Election Date: 1/1/0001" and no table -- and
#: that is an absence, not a schema change. Above this, a missing header IS drift.
MIN_TABLE_ROWS = 10

#: A filename that is certainly NOT the general election. Ordering only.
_PRIMARY_HINT = re.compile(r"primar|runoff|ppp|special|\bdem\b|\brep\b", re.I)
#: Which of the two reports a filename is. Stable in every cycle since 2022.
_ABSENTEE_FILE = re.compile(r"absentee", re.I)
_EARLY_FILE = re.compile(r"early.?voting", re.I)

_DATE = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b")
_RACE_LINE = re.compile(r"^race\s*:?\s*(.+)$", re.I)
_PARTY_LINE = re.compile(r"^political party\s*:?\s*(.+)$", re.I)
_COUNTY_CODE = re.compile(r"^\d+\s*-\s*")


def _norm(value) -> str:
    return " ".join(str(value or "").split()).lower()


def _int(value) -> int | None:
    """A count, or None when the cell is blank ("South Carolina did not report")."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise SchemaDrift(f"SC: {value!r} is not a count")
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", "")
    if not text or text in {"-", "--", "N/A", "n/a"}:
        return None
    try:
        return int(float(text))
    except ValueError as exc:
        raise SchemaDrift(f"SC: {value!r} is not a count") from exc


def _add(*values: int | None) -> int | None:
    """Sum the reported parts, or None if no part was reported at all."""
    known = [v for v in values if v is not None]
    return sum(known) if known else None


def early_window(cycle: int) -> tuple[date, date]:
    """The statutory in-person early-voting window for `cycle`."""
    end = election_date(cycle)
    return end - timedelta(days=EV_OPEN_OFFSET), end - timedelta(days=EV_CLOSE_OFFSET)


def _parse_dates(text: str) -> set[date]:
    out: set[date] = set()
    for month, day, year in _DATE.findall(text or ""):
        try:
            out.add(date(int(year), int(month), int(day)))
        except ValueError:
            continue
    return out


def _cell_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    found = _parse_dates(str(value or ""))
    return next(iter(found)) if len(found) == 1 else None


def _open(body: bytes, what: str):
    if not looks_like_xlsx(body):
        raise SourceError(f"SC: {what} was not an xlsx")
    try:
        return openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- openpyxl raises a zoo of types
        raise SourceError(f"SC: could not open {what}: {exc}") from exc


# --------------------------------------------------------------------------
# The absentee (mail) workbook
# --------------------------------------------------------------------------
def _sheet_layout(rows: list[tuple]) -> tuple[int, int, int] | None:
    """(header row index, county column, first Ballots column) for one worksheet.

    Returns None for a sheet with no "County" header at all -- the 2026 report
    carries cover-style sheets -- and raises SchemaDrift for a sheet that has one
    but whose measure banners do not line up, because that is the case where a
    positional read would publish the wrong column.
    """
    for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        county_at = next((j for j, c in enumerate(row) if _norm(c) == COUNTY), None)
        if county_at is None:
            continue
        if i < 2:
            raise SchemaDrift("SC: absentee header has no measure banners above it")

        groups = {_norm(c): j for j, c in enumerate(rows[i - 2]) if _norm(c)}
        if GROUP_TOTAL not in groups:
            raise SchemaDrift(f"SC: absentee report has no {GROUP_TOTAL!r} column group")
        start = groups[GROUP_TOTAL]
        # Everything we read must sit inside the Total group, i.e. left of the
        # next group banner (Non-UOCAVA). Reading past it would silently publish
        # the domestic-only sub-total as the whole.
        limit = min((j for j in groups.values() if j > start), default=len(rows[i - 2]))

        subs = [(j, _norm(c)) for j, c in enumerate(rows[i - 1]) if _norm(c)]
        ballots_at = next((j for j, name in subs
                           if name == SUB_BALLOTS and start <= j < limit), None)
        if ballots_at is None:
            raise SchemaDrift(
                f"SC: absentee report has no {SUB_BALLOTS!r} block inside {GROUP_TOTAL!r}")

        labels = [_norm(row[j]) if j < len(row) else ""
                  for j in range(ballots_at, ballots_at + len(BALLOT_MEASURES))]
        if tuple(labels) != BALLOT_MEASURES:
            raise SchemaDrift(f"SC: absentee ballot columns are {labels}, "
                              f"expected {list(BALLOT_MEASURES)}")
        return i, county_at, ballots_at
    return None


def _sheet_banner(rows: list[tuple], header_at: int) -> str:
    return " ".join(str(c) for row in rows[:header_at] for c in row
                    if isinstance(c, str) and c.strip())


def _labelled(rows: list[tuple], header_at: int, pattern: re.Pattern) -> str | None:
    """The value of a `Race:` / `Political Party:` line above the header, if any."""
    for row in rows[:header_at]:
        for cell in row:
            if not isinstance(cell, str):
                continue
            m = pattern.match(" ".join(cell.split()))
            if m:
                return m.group(1).strip()
    return None


def race_bucket(raw: str) -> str:
    """Map a South Carolina worksheet's race label onto a canonical bucket."""
    bucket = _race(raw) or EXTRA_RACE.get(_norm(raw))
    if bucket is None:
        raise SchemaDrift(f"SC: unrecognised race label {raw!r}")
    return bucket


def check_election(rows: list[tuple], header_at: int, cycle: int) -> None:
    """Refuse a workbook that is not this cycle's GENERAL election.

    South Carolina posts the same report for its primaries, its runoffs and its
    specials, under filenames that differ only by a word. The banner above the
    header names the election, and that is what decides.
    """
    banner = _sheet_banner(rows, header_at)
    if "primary" in banner.lower():
        raise NotYetPublished(f"SC: {banner.strip()[:90]!r} is a primary, not the general")
    if election_date(cycle) not in _parse_dates(banner):
        raise NotYetPublished(
            f"SC: workbook names no {election_date(cycle).isoformat()} election "
            f"({banner.strip()[:90]!r})"
        )


def parse_absentee(body: bytes, cycle: int, day: date, *, with_race: bool) -> FetchResult:
    """Parse one absentee workbook: county mail rows, plus race if in scope.

    `ballots_total` is set to mail returned here; the caller folds in-person into
    it when the early-voting file says what in-person was.
    """
    book = _open(body, "absentee report")
    result = FetchResult()
    seen_race: dict[str, str] = {}
    by_bucket: dict[str, int | None] = {}
    checked = False
    widest = 0

    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        widest = max(widest, sum(1 for r in rows if any(c not in (None, "") for c in r)))
        layout = _sheet_layout(rows)
        if layout is None:
            continue
        header_at, county_at, ballots_at = layout
        if not checked:
            check_election(rows, header_at, cycle)
            checked = True

        # A 2026-style sheet filtered to one party's primary ballot is not a
        # statewide universe and must never be read as one.
        party = _labelled(rows, header_at, _PARTY_LINE)
        if party is not None and _norm(party) not in {"all", ""}:
            continue

        label = _labelled(rows, header_at, _RACE_LINE)
        if label is None:
            raise SchemaDrift(f"SC: worksheet {sheet.title!r} names no race filter")
        if _norm(label) == "all":
            if result.county_rows:
                raise SchemaDrift(
                    f"SC: worksheet {sheet.title!r} is a second unfiltered 'Race ALL' "
                    f"table; we cannot tell which one is the state")
            _emit_absentee(result, rows, header_at, county_at, ballots_at, cycle, day)
            continue
        if not with_race:
            continue

        if _norm(label) in seen_race:
            raise SchemaDrift(
                f"SC: worksheets {seen_race[_norm(label)]!r} and {sheet.title!r} are "
                f"both the {label!r} table")
        seen_race[_norm(label)] = sheet.title
        total = _statewide_returned(rows, header_at, county_at, ballots_at)
        if total is not None:
            # `normalize.RACE_BUCKETS` has no multiracial bucket, so South
            # Carolina's "Multiple" and "Other" worksheets are both `other` and
            # have to be ADDED. Raising on the collision would leave the state
            # permanently broken; dropping one would lose real ballots.
            by_bucket[race_bucket(label)] = _add(by_bucket.get(race_bucket(label)), total)

    for bucket, total in sorted(by_bucket.items()):
        result.demo_rows.append(DemoDay(
            cycle=cycle, state="SC", day=day,
            dimension="race", bucket=bucket, ballots_total=total,
        ))
    if not checked:
        if widest < MIN_TABLE_ROWS:
            raise NotYetPublished(
                "SC: the posted absentee workbook is the SEC's empty shell -- "
                "a banner with no table under it")
        raise SchemaDrift("SC: absentee report has no worksheet with a 'County' header")
    if not result.county_rows:
        raise SchemaDrift("SC: absentee report produced no county rows")
    return result


def _data_rows(rows: list[tuple], header_at: int, county_at: int):
    for row in rows[header_at + 1:]:
        if county_at >= len(row):
            continue
        name = " ".join(str(row[county_at] or "").split())
        if name:
            yield name, row


def _statewide_returned(rows, header_at, county_at, ballots_at) -> int | None:
    for name, row in _data_rows(rows, header_at, county_at):
        if name.lower() == TOTAL_LABEL:
            i = ballots_at + 3
            return _int(row[i]) if i < len(row) else None
    return None


def _emit_absentee(result, rows, header_at, county_at, ballots_at, cycle, day) -> None:
    unknown: list[str] = []
    for name, row in _data_rows(rows, header_at, county_at):
        def cell(offset):
            i = ballots_at + offset
            return _int(row[i]) if i < len(row) else None

        issued, returned = cell(0), cell(3)
        if name.lower() == TOTAL_LABEL:
            # South Carolina's own statewide line, so our headline is theirs.
            result.state_rows.append(StateDay(
                cycle=cycle, state="SC", day=day,
                ballots_total=returned,
                mail_requested=issued,
                mail_returned=returned,
                # South Carolina does not register voters by party.
                party_dem=None, party_rep=None, party_npa=None, party_oth=None,
            ))
            continue

        hit = _fips.lookup("SC", _COUNTY_CODE.sub("", name))
        if hit is None:
            unknown.append(name)
            continue
        fips, canonical = hit
        result.county_rows.append(CountyDay(
            cycle=cycle, state="SC", county_fips=fips, day=day,
            county_name=canonical,
            ballots_total=returned,
            mail_returned=returned,
            party_dem=None, party_rep=None, party_npa=None, party_oth=None,
        ))

    if unknown:
        # South Carolina has exactly 46 counties and they do not change.
        raise SchemaDrift(f"SC: unrecognised county names {sorted(set(unknown))[:5]}")


# --------------------------------------------------------------------------
# The in-person early-voting workbook
# --------------------------------------------------------------------------
EV_HEADER = ("jurisdiction", "election name", "poll place")
EV_TOTAL_COLUMN = "total voters"
EV_STATEWIDE = "south carolina"


class EarlyVoting:
    """One parsed early-voting workbook: the in-person curve, by county and day."""

    __slots__ = ("report_day", "days", "statewide", "counties", "names")

    def __init__(self, report_day: date, days: list[date]) -> None:
        self.report_day = report_day
        self.days = days
        self.statewide: dict[date, int] = {}
        self.counties: dict[str, dict[date, int]] = {}
        self.names: dict[str, str] = {}

    def through(self, table: dict[date, int], day: date) -> int:
        return sum(n for d, n in table.items() if d <= day)


def parse_early(body: bytes, cycle: int) -> EarlyVoting:
    """Parse the early-voting workbook into per-county, per-day in-person counts."""
    book = _open(body, "early-voting report")
    rows = list(book.worksheets[0].iter_rows(values_only=True))

    header_at = next((i for i, row in enumerate(rows[:HEADER_SEARCH_ROWS])
                      if row and _norm(row[0]) == EV_HEADER[0]), None)
    if header_at is None:
        raise SchemaDrift("SC: early-voting report has no 'Jurisdiction' header row")
    header = rows[header_at]
    if tuple(_norm(c) for c in header[:3]) != EV_HEADER:
        raise SchemaDrift(f"SC: early-voting header starts {[_norm(c) for c in header[:3]]}")

    day_columns: list[tuple[int, date]] = []
    total_at = None
    for j, cell in enumerate(header[3:], start=3):
        if _norm(cell) == EV_TOTAL_COLUMN:
            total_at = j
            continue
        parsed = _cell_date(cell)
        if parsed is not None:
            day_columns.append((j, parsed))
        elif _norm(cell):
            raise SchemaDrift(f"SC: early-voting column {cell!r} is neither a date "
                              f"nor {EV_TOTAL_COLUMN!r}")
    if not day_columns or total_at is None:
        raise SchemaDrift("SC: early-voting report has no date columns or no total column")

    # The banner above the header carries the election it is for and the day it
    # was run; both are labelled, so both are read by label.
    banner = {_norm(c): (i, j) for i, row in enumerate(rows[:header_at])
              for j, c in enumerate(row) if _norm(c)}
    election_at = banner.get("election date")
    report_at = banner.get("report date")
    if election_at is None or report_at is None:
        raise SchemaDrift("SC: early-voting report has no Election Date / Report Date")
    named = _cell_date(rows[election_at[0] + 1][election_at[1]])
    if named != election_date(cycle):
        raise NotYetPublished(
            f"SC: early-voting report is for {named}, not {election_date(cycle)}")
    report_day = _cell_date(rows[report_at[0] + 1][report_at[1]])
    if report_day is None:
        raise SchemaDrift("SC: early-voting report has no readable report date")

    ev = EarlyVoting(report_day, [d for _, d in day_columns])
    unknown: list[str] = []
    for row in rows[header_at + 1:]:
        name = " ".join(str(row[0] or "").split()) if row else ""
        if not name:
            # A polling-place row, or a blank separator. Polling places are the
            # county's own subtotals; counting them would double the state.
            continue
        # A blank day cell in this file is "nobody voted here that day", not
        # "not reported": the report is generated from the state's own system and
        # only lists days the county was open.
        counts = {d: (_int(row[j]) or 0) if j < len(row) else 0 for j, d in day_columns}
        stated = _int(row[total_at]) if total_at < len(row) else None
        if stated is not None and stated != sum(counts.values()):
            # The only way this happens is a day column we failed to see, which
            # would understate every county in the state.
            raise SchemaDrift(
                f"SC: {name} early-voting days sum to {sum(counts.values())} but the "
                f"report says {stated}")

        if _norm(name) == EV_STATEWIDE:
            ev.statewide = counts
            continue
        hit = _fips.lookup("SC", name)
        if hit is None:
            unknown.append(name)
            continue
        fips, canonical = hit
        ev.counties[fips] = counts
        ev.names[fips] = canonical

    if unknown:
        raise SchemaDrift(f"SC: unrecognised early-voting jurisdictions {sorted(set(unknown))[:5]}")
    if not ev.statewide or not ev.counties:
        raise SchemaDrift("SC: early-voting report produced no rows")
    return ev


# --------------------------------------------------------------------------
# Folding the two together
# --------------------------------------------------------------------------
def combine(mail: FetchResult, ev: EarlyVoting | None, cycle: int, day: date,
            *, with_curve: bool = True) -> FetchResult:
    """Add in-person to the absentee snapshot, and emit the in-person curve.

    `mail` is folded into IN PLACE and returned.

    The absentee snapshot's day is the anchor: on it, `ballots_total` is mail
    returned plus in-person cast through the last day the early-voting file
    covers. Earlier days get an in-person row with `ballots_total` blank, because
    what mail had come back on those days is genuinely not known from these two
    files -- blank, never a total that quietly omits the mail. See THE BLANK RULE.
    """
    if ev is None or (day - ev.report_day).days > EV_MAX_SKEW_DAYS:
        if ev is not None:
            log.warning("SC: early-voting report %s is too far behind the absentee "
                        "snapshot %s to combine", ev.report_day, day)
        if day >= early_window(cycle)[0]:
            # In-person voting is running and we do not know how much of it has
            # happened, so the day's TOTAL is genuinely unknown. Leaving it as
            # mail alone would publish a number an order of magnitude short.
            for row in mail.state_rows + mail.county_rows:
                row.ballots_total = None
        return mail

    anchor_days = [d for d in ev.days if d <= day]
    if not anchor_days:
        return mail
    anchor = max(anchor_days)

    for row in mail.state_rows:
        row.inperson = ev.through(ev.statewide, anchor)
        row.ballots_total = _add(row.mail_returned, row.inperson)
    for row in mail.county_rows:
        counts = ev.counties.get(row.county_fips)
        if counts is None:
            # In-person is open and we do not know this county's share of it, so
            # its total is unknown -- not "its mail is its total".
            row.ballots_total = None
            continue
        row.inperson = ev.through(counts, anchor)
        row.ballots_total = _add(row.mail_returned, row.inperson)

    if not with_curve:
        return mail

    for d in anchor_days:
        if d == day:
            continue  # already covered by the anchor row above
        mail.state_rows.append(StateDay(
            cycle=cycle, state="SC", day=d,
            ballots_total=None, inperson=ev.through(ev.statewide, d),
            mail_requested=None, mail_returned=None,
            party_dem=None, party_rep=None, party_npa=None, party_oth=None,
        ))
        for fips, counts in sorted(ev.counties.items()):
            mail.county_rows.append(CountyDay(
                cycle=cycle, state="SC", county_fips=fips, day=d,
                county_name=ev.names[fips],
                ballots_total=None, inperson=ev.through(counts, d),
                mail_returned=None,
                party_dem=None, party_rep=None, party_npa=None, party_oth=None,
            ))
    return mail


# --------------------------------------------------------------------------
class SCScraper(Adapter):
    """Tier 1 for South Carolina: the SEC's absentee and early-voting workbooks."""

    state = "SC"
    name = "sc-sec"
    tier = TIER_SCRAPER

    # ------------------------------------------------------------------
    def _media(self, cycle: int) -> list[dict]:
        """Every workbook the SEC's media library lists for this cycle's year."""
        found: dict[str, dict] = {}
        for term in SEARCH_TERMS:
            slug = term.replace(" ", "_")
            try:
                body = get(MEDIA, state="SC", filename=f"media_{cycle}_{slug}.json",
                           params={"search": term, "per_page": 100, "orderby": "date",
                                   "order": "desc", "_fields": "date,source_url"},
                           min_bytes=2)
            except Missing:
                continue
            try:
                listing = json.loads(body.decode("utf-8", errors="replace"))
            except ValueError as exc:
                raise SourceError(f"SC: media index for {term!r} was not JSON: {exc}") from exc
            if not isinstance(listing, list):
                raise SourceError(f"SC: media index for {term!r} was not a list")
            for item in listing:
                url = str(item.get("source_url") or "")
                stamp = str(item.get("date") or "")[:10]
                if not url.lower().endswith(".xlsx") or not stamp.startswith(str(cycle)):
                    continue
                try:
                    uploaded = date.fromisoformat(stamp)
                except ValueError:
                    continue
                found[url] = {"url": url, "day": uploaded,
                              "at": str(item.get("date") or ""),
                              "file": url.rsplit("/", 1)[-1]}
        # Newest first, by full timestamp: South Carolina posts several snapshots
        # on the last few days ("…11-4-9-am", "…11-4-4-pm") and the latest wins.
        return sorted(found.values(), key=lambda m: m["at"], reverse=True)

    @staticmethod
    def _candidates(media: list[dict], wanted: re.Pattern) -> tuple[list[dict], list[dict]]:
        """(files that look like the general, everything else of that report)."""
        matched = [c for c in media if wanted.search(c["file"])]
        likely = [c for c in matched if not _PRIMARY_HINT.search(c["file"])]
        return likely, [c for c in matched if c not in likely]

    @staticmethod
    def _rank(candidates: list[dict], wanted: re.Pattern) -> list[dict]:
        """The files of one report, newest first, obvious primaries last.

        The report words ("Absentee", "Early Voting") have been stable in every
        filename since 2022 and select WHICH report; everything else in the name
        is only a hint about WHICH ELECTION, so a primary-looking name is pushed
        to the back rather than dropped -- if South Carolina renames the
        general's file we must still find it, and the workbook's own banner is
        what actually decides.
        """
        likely, rest = SCScraper._candidates(candidates, wanted)
        return likely + rest

    def _download(self, item: dict, *, use_cache: bool) -> bytes | None:
        try:
            body = get(item["url"], state="SC", filename=item["file"],
                       use_cache=use_cache, min_bytes=4096)
        except Missing:
            return None
        return body if looks_like_xlsx(body) else None

    # ------------------------------------------------------------------
    def _load_absentee(self, cycle: int, as_of: date, media: list[dict],
                       *, use_cache: bool) -> tuple[date, FetchResult]:
        window = early_window(cycle)
        fresh = [m for m in media if m["day"] <= as_of]
        problems: list[str] = []
        for item in self._rank(fresh, _ABSENTEE_FILE)[:MAX_PROBE]:
            body = self._download(item, use_cache=use_cache)
            if body is None:
                problems.append(f"{item['file']} did not download as an xlsx")
                continue
            try:
                return item["day"], parse_absentee(
                    body, cycle, item["day"], with_race=item["day"] < window[0])
            except NotYetPublished as exc:
                if str(exc) not in problems:
                    problems.append(str(exc))
        raise NotYetPublished(
            f"SC: no {cycle} general-election absentee workbook posted by "
            f"{as_of.isoformat()} ({'; '.join(problems) or 'nothing in the media index'})"
        )

    def _load_early(self, cycle: int, as_of: date, media: list[dict],
                    *, use_cache: bool) -> EarlyVoting | None:
        fresh = [m for m in media if m["day"] <= as_of]
        for item in self._rank(fresh, _EARLY_FILE)[:MAX_PROBE]:
            body = self._download(item, use_cache=use_cache)
            if body is None:
                continue
            try:
                return parse_early(body, cycle)
            except NotYetPublished as exc:
                log.debug("SC: %s", exc)
        return None

    def _run(self, cycle: int, as_of: date, *, use_cache: bool) -> FetchResult:
        media = self._media(cycle)
        if not media:
            raise NotYetPublished(f"SC: the media library lists no {cycle} workbooks yet")
        day, mail = self._load_absentee(cycle, as_of, media, use_cache=use_cache)
        ev = self._load_early(cycle, as_of, media, use_cache=use_cache)
        if ev is None and day >= early_window(cycle)[0]:
            # In-person voting is open and is the great majority of South
            # Carolina's early vote, so an absentee-only answer would understate
            # the headline by an order of magnitude. Fall through a tier instead.
            raise SourceError(
                f"SC: in-person early voting opened {early_window(cycle)[0].isoformat()} "
                f"but no early-voting workbook could be read"
            )
        return combine(mail, ev, cycle, day)

    # ------------------------------------------------------------------
    def fetch(self, cycle: int, as_of: date) -> FetchResult:
        return self._run(cycle, as_of, use_cache=False)

    def fetch_history(self, cycle: int) -> FetchResult:
        """Every archived daily snapshot for a past cycle.

        South Carolina's media library keeps every dated absentee workbook, so the
        archive really is a daily series; and the LAST early-voting workbook
        carries every day of in-person voting in its columns, so one download
        supplies the in-person cumulative for each of those snapshot days. That
        makes a past cycle a real curve rather than a single final.

        Only files whose name does not look like a primary are walked here -- a
        cycle's archive holds dozens of presidential-preference, primary and
        runoff reports, and each one is a download that would only be rejected.
        """
        if cycle >= datetime.now().year:
            raise NotYetPublished(f"SC: {cycle} is not an archived cycle")
        media = self._media(cycle)
        if not media:
            raise NotYetPublished(f"SC: the media library lists no {cycle} workbooks")

        as_of = election_date(cycle)
        opens = early_window(cycle)[0]
        ev = self._load_early(cycle, as_of, media, use_cache=True)
        candidates, _ = self._candidates([m for m in media if m["day"] <= as_of],
                                         _ABSENTEE_FILE)
        combined = FetchResult()
        seen: set[date] = set()
        for item in candidates:
            if item["day"] in seen:
                continue          # an earlier snapshot from a day we already have
            body = self._download(item, use_cache=True)
            if body is None:
                continue
            try:
                snapshot = parse_absentee(body, cycle, item["day"],
                                          with_race=item["day"] < opens)
            except NotYetPublished as exc:
                log.debug("SC: %s", exc)
                continue
            seen.add(item["day"])
            combined.extend(combine(snapshot, ev, cycle, item["day"], with_curve=False))

        if not combined:
            raise NotYetPublished(f"SC: no archived {cycle} general-election workbooks")
        return combined
